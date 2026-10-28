import csv
from datetime import date, datetime, time, timezone as datetime_timezone
from io import StringIO

from django.test import RequestFactory, SimpleTestCase, TestCase, override_settings

from smartbids.utils.licitacion_utils import estado_licitacion, escribir_fila_csv
from smartbids.models.procurement import Licitacion
from smartbids.templatetags.montos import miles_chilenos


@override_settings(TIME_ZONE='UTC')
class EstadoLicitacionTests(SimpleTestCase):
    ahora = datetime(2026, 10, 13, 12, 0, tzinfo=datetime_timezone.utc)

    def test_calcula_periodo_de_preguntas(self):
        licitacion = Licitacion(
            lic_fecha_cierre=date(2026, 10, 20),
            lic_fecha_preguntas=datetime(2026, 10, 13, 11, 0, tzinfo=datetime_timezone.utc),
            lic_fecha_respuestas=datetime(2026, 10, 14, 11, 0, tzinfo=datetime_timezone.utc),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora), ('preguntas', 'Período de preguntas'))

    def test_calcula_recepcion_de_ofertas(self):
        licitacion = Licitacion(
            lic_fecha_cierre=date(2026, 10, 20),
            lic_fecha_respuestas=datetime(2026, 10, 13, 11, 0, tzinfo=datetime_timezone.utc),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora)[0], 'ofertas')

    def test_acepta_fechas_legacy_sin_zona_horaria(self):
        licitacion = Licitacion(
            lic_fecha_cierre=date(2026, 10, 20),
            lic_fecha_preguntas=datetime(2026, 10, 13, 11, 0),
            lic_fecha_respuestas=datetime(2026, 10, 14, 11, 0),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora)[0], 'preguntas')

    def test_cierre_tiene_precedencia_sobre_otros_hitos(self):
        licitacion = Licitacion(
            lic_fecha_cierre=date(2026, 10, 12),
            lic_fecha_respuestas=datetime(2026, 10, 13, 11, 0, tzinfo=datetime_timezone.utc),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora)[0], 'cerrada')

    def test_estado_inicial_si_no_hay_hitos_vencidos(self):
        licitacion = Licitacion(
            lic_fecha_public=date(2026, 10, 13),
            lic_fecha_cierre=date(2026, 10, 20),
            lic_hora_cierre=time(17, 0),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora)[0], 'publicada')

    def test_exportacion_csv_protege_formulas_de_hoja_de_calculo(self):
        output = StringIO()
        writer = csv.writer(output)
        escribir_fila_csv(writer, ('=1+1', 'Texto normal'))
        self.assertEqual(next(csv.reader(StringIO(output.getvalue()))), ["'=1+1", 'Texto normal'])

    def test_exportacion_excel_genera_archivo_y_protege_formulas(self):
        from io import BytesIO
        import openpyxl
        from smartbids.utils.licitacion_utils import generar_excel_response

        response = generar_excel_response(
            'test.xlsx',
            ('Columna 1', 'Columna 2'),
            [('=SUM(A1:A2)', 'Texto simple'), ('@link', 123)]
        )
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('test.xlsx', response['Content-Disposition'])

        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, 'Columna 1')
        self.assertEqual(ws.cell(row=2, column=1).value, "'=SUM(A1:A2)")
        self.assertEqual(ws.cell(row=2, column=2).value, 'Texto simple')
        self.assertEqual(ws.cell(row=3, column=1).value, "'@link")
        self.assertEqual(ws.cell(row=3, column=2).value, 123)


class FormatoMontosTests(SimpleTestCase):
    def test_formatea_montos_con_separador_de_miles_chileno(self):
        self.assertEqual(miles_chilenos('125000000'), '125.000.000')

    def test_redondea_montos_sin_decimales(self):
        self.assertEqual(miles_chilenos('125000000.50'), '125.000.001')


class PreferenciasTests(SimpleTestCase):
    def test_normalizar_cadena_simple(self):
        from smartbids.views.perfil_views import normalizar_cadena_simple
        self.assertEqual(normalizar_cadena_simple("Mantención"), "mantencion")
        self.assertEqual(normalizar_cadena_simple("¡Cloud & Seguridad!"), "cloud seguridad")
        self.assertEqual(normalizar_cadena_simple(""), "")

    def test_nlp_sanitizar_caracteres_especiales(self):
        from smartbids.utils.palabras_utils import sanitizar_caracteres_especiales
        texto_con_control = "Servidor\x00\x08 IBM\x1F Cloud\ue000"
        self.assertEqual(sanitizar_caracteres_especiales(texto_con_control), "Servidor IBM Cloud")

    def test_nlp_filtrar_stopwords(self):
        from smartbids.utils.palabras_utils import filtrar_STOPWORDS
        sw = {"de", "y", "el", "la", "en"}
        texto = "sistema de gestion y control en la nube"
        self.assertEqual(filtrar_STOPWORDS(texto, sw), "sistema gestion control nube")

    def test_nlp_obtener_palabras_unicas_ngramas(self):
        from smartbids.utils.palabras_utils import obtener_palabras_unicas
        texto = "sistema riesgo"
        # npal = 1 (unigramas)
        self.assertEqual(obtener_palabras_unicas(texto, npal=1), {"sistema", "riesgo"})
        # npal = 2 (unigramas y bigramas)
        self.assertEqual(obtener_palabras_unicas(texto, npal=2), {"sistema", "riesgo", "sistema riesgo"})
        # npal = 3 (unigramas, bigramas y trigramas)
        texto_tri = "mesa ayuda tecnica"
        tokens_tri = obtener_palabras_unicas(texto_tri, npal=3)
        self.assertIn("mesa", tokens_tri)
        self.assertIn("ayuda", tokens_tri)
        self.assertIn("tecnica", tokens_tri)
        self.assertIn("mesa ayuda", tokens_tri)
        self.assertIn("ayuda tecnica", tokens_tri)
        self.assertIn("mesa ayuda tecnica", tokens_tri)

    def test_nlp_normalizar_texto_completo(self):
        from smartbids.utils.palabras_utils import normalizar_texto
        sw = {"de", "para"}
        # Tildes, números aislados, signos de puntuación y stopwords
        texto = "Mantención de Computadores +2026 para Diseño Web & Redes!"
        norm = normalizar_texto(texto, stopwords_set=sw)
        self.assertEqual(norm, "mantencion computadores diseno web redes")

    def test_nlp_procesar_palabras_usuario_dual(self):
        from unittest.mock import patch
        from smartbids.utils.palabras_utils import procesar_palabras_usuario
        with patch('smartbids.utils.palabras_utils.leer_STOPWORDS', return_value={"de", "y", "el"}):
            entrada = ["Computadora", "Sistema de Riesgo", "Auditoría ISO 9001"]
            orig, norm = procesar_palabras_usuario(entrada, npal=3)
            # Escritura original del usuario intacta
            self.assertEqual(orig, ["Computadora", "Sistema de Riesgo", "Auditoría ISO 9001"])
            # Conjunto normalizado con unigramas y n-gramas
            self.assertIn("computadora", norm)
            self.assertIn("sistema", norm)
            self.assertIn("riesgo", norm)
            self.assertIn("sistema riesgo", norm)
            self.assertIn("auditoria", norm)
            self.assertIn("iso", norm)
            self.assertIn("auditoria iso", norm)
            # El número aislado 9001 y la stopword 'de' no deben estar
            self.assertNotIn("9001", norm)
            self.assertNotIn("de", norm)

    def test_actualizar_preferencias_api_guarda_campos_duales(self):
        import json
        from unittest.mock import patch, MagicMock
        from django.test import RequestFactory
        from smartbids.views.perfil_views import actualizar_preferencias_suscriptor

        rf = RequestFactory()
        body = json.dumps({
            'uid': 'test-firebase-uid',
            'pref_palabras_claves': ['Computadora', 'Sistema de Riesgo']
        }).encode('utf-8')
        req = rf.post('/api/actualizar-preferencias/', data=body, content_type='application/json')

        mock_suscriptor = MagicMock()
        mock_pref = MagicMock()
        mock_pref.pref_palabras_claves = []
        mock_pref.pref_palabras_claves_norm = []

        with patch('smartbids.views.perfil_views.Suscriptor.objects.filter') as mock_filter, \
             patch('smartbids.views.perfil_views.validar_autorizacion_suscriptor', return_value=(True, '')), \
             patch('smartbids.views.perfil_views.Preferencia.objects.get_or_create', return_value=(mock_pref, False)), \
             patch('smartbids.utils.palabras_utils.leer_STOPWORDS', return_value={"de"}):
            
            mock_filter.return_value.first.return_value = mock_suscriptor
            resp = actualizar_preferencias_suscriptor(req)
            self.assertEqual(resp.status_code, 200)
            data = json.loads(resp.content.decode('utf-8'))
            self.assertEqual(data['status'], 'ok')
            # Verifica que la respuesta mantenga las palabras del usuario intactas
            self.assertEqual(data['palabras_sanitizadas'], ['Computadora', 'Sistema de Riesgo'])
            # Verifica que pref_palabras_claves contenga la escritura del usuario
            self.assertEqual(mock_pref.pref_palabras_claves, ['Computadora', 'Sistema de Riesgo'])
            # Verifica que pref_palabras_claves_norm contenga los tokens normalizados
            self.assertIn('computadora', mock_pref.pref_palabras_claves_norm)
            self.assertIn('sistema riesgo', mock_pref.pref_palabras_claves_norm)
            self.assertNotIn('de', mock_pref.pref_palabras_claves_norm)


class VistasExportacionTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        self.rf = RequestFactory()

    def test_exportar_licitaciones_csv_por_defecto(self):
        from unittest.mock import patch
        from smartbids.views.licitaciones_views import exportar_licitaciones_view

        with patch('smartbids.views.licitaciones_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.views.licitaciones_views.filtrar_licitaciones', return_value=([], 0)):
            req = self.rf.get('/mis-licitaciones/exportar/')
            resp = exportar_licitaciones_view(req)
            self.assertEqual(resp['Content-Type'], 'text/csv; charset=utf-8')
            self.assertIn('licitaciones-smartbids.csv', resp['Content-Disposition'])

    def test_exportar_licitaciones_excel(self):
        from unittest.mock import patch
        from smartbids.views.licitaciones_views import exportar_licitaciones_view

        with patch('smartbids.views.licitaciones_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.views.licitaciones_views.filtrar_licitaciones', return_value=([], 0)):
            req = self.rf.get('/mis-licitaciones/exportar/?formato=xlsx')
            resp = exportar_licitaciones_view(req)
            self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            self.assertIn('licitaciones-smartbids.xlsx', resp['Content-Disposition'])

    def test_exportar_dashboard_excel(self):
        from unittest.mock import patch
        from smartbids.views.page_views import exportar_dashboard_view

        req = self.rf.get('/dashboard/exportar/?formato=xlsx&gran=mes&periodo=2024-05')
        resp = exportar_dashboard_view(req)
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('dashboard-tablas-2024-05.xlsx', resp['Content-Disposition'])

    def test_exportar_dashboard_csv(self):
        from smartbids.views.page_views import exportar_dashboard_view

        req = self.rf.get('/dashboard/exportar/?formato=csv&gran=mes&periodo=2024-05')
        resp = exportar_dashboard_view(req)
        self.assertEqual(resp['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('dashboard-tablas-2024-05.csv', resp['Content-Disposition'])
        self.assertIn('=== RESUMEN KPIS ===', resp.content.decode('utf-8'))



class DashboardDWTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def test_dashboard_page_render(self):
        from unittest.mock import patch
        from smartbids.views.dashboard_views import dashboard_view
        req = self.rf.get('/dashboard/')
        with patch('smartbids.views.dashboard_views.render') as mock_render:
            mock_render.return_value = 'ok'
            resp = dashboard_view(req)
            self.assertEqual(resp, 'ok')
            mock_render.assert_called_once_with(req, 'smartbids/dashboard.html')

    def test_dashboard_periodo_invalido(self):
        from smartbids.views.dashboard_views import dashboard_kpis
        req = self.rf.get('/dashboard/kpis/?gran=mensual&periodo=invalido')
        resp = dashboard_kpis(req)
        self.assertEqual(resp.status_code, 400)


class ReenvioActivacionTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def test_metodo_no_permitido(self):
        from smartbids.views.auth_views import reenviar_correo_activacion
        req = self.rf.get('/api/reenviar-correo-activacion/')
        resp = reenviar_correo_activacion(req)
        self.assertEqual(resp.status_code, 405)

    def test_correo_requerido(self):
        from smartbids.views.auth_views import reenviar_correo_activacion
        req = self.rf.post('/api/reenviar-correo-activacion/', data=b'{}', content_type='application/json')
        resp = reenviar_correo_activacion(req)
        self.assertEqual(resp.status_code, 400)

    def test_usuario_no_encontrado(self):
        from unittest.mock import patch, MagicMock
        from firebase_admin import auth as fb_auth
        from smartbids.views.auth_views import reenviar_correo_activacion

        mock_app = MagicMock()
        with patch('smartbids.views.auth_views._firebase_admin_app', return_value=mock_app), \
             patch('firebase_admin.auth.get_user_by_email', side_effect=fb_auth.UserNotFoundError('No user')):
            req = self.rf.post(
                '/api/reenviar-correo-activacion/',
                data=b'{"email": "inexistente@smartbids.cl"}',
                content_type='application/json'
            )
            resp = reenviar_correo_activacion(req)
            self.assertEqual(resp.status_code, 404)

    def test_usuario_ya_activado(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.auth_views import reenviar_correo_activacion

        mock_user = MagicMock(email_verified=True)
        with patch('smartbids.views.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', return_value=mock_user):
            req = self.rf.post(
                '/api/reenviar-correo-activacion/',
                data=b'{"email": "activo@smartbids.cl"}',
                content_type='application/json'
            )
            resp = reenviar_correo_activacion(req)
            self.assertEqual(resp.status_code, 400)

    def test_reenvio_exitoso_y_rate_limit(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.auth_views import reenviar_correo_activacion

        mock_user = MagicMock(email_verified=False)
        with patch('smartbids.views.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', return_value=mock_user), \
             patch('smartbids.views.auth_views._enviar_email_activacion') as mock_envio:
            req = self.rf.post(
                '/api/reenviar-correo-activacion/',
                data=b'{"email": "pendiente@smartbids.cl"}',
                content_type='application/json'
            )
            resp = reenviar_correo_activacion(req)
            self.assertEqual(resp.status_code, 200)
            mock_envio.assert_called_once_with('pendiente@smartbids.cl', app=mock_envio.call_args.kwargs.get('app'), request=req)

            # Segundo intento inmediato debe devolver 429
            req2 = self.rf.post(
                '/api/reenviar-correo-activacion/',
                data=b'{"email": "pendiente@smartbids.cl"}',
                content_type='application/json'
            )
            resp2 = reenviar_correo_activacion(req2)
            self.assertEqual(resp2.status_code, 429)

    def test_correo_malformado_admin(self):
        import json
        from smartbids.views.auth_views import reenviar_correo_activacion

        req = self.rf.post(
            '/api/reenviar-correo-activacion/',
            data=b'{"email": "admin"}',
            content_type='application/json'
        )
        resp = reenviar_correo_activacion(req)
        self.assertEqual(resp.status_code, 400)
        data = json.loads(resp.content.decode('utf-8'))
        self.assertEqual(data.get('status'), 'error')
        self.assertIn('formato del correo ingresado no es válido', data.get('mensaje', ''))
        self.assertNotIn('Malformed email address', data.get('mensaje', ''))

    def test_correo_malformado_en_enviar_codigo_login(self):
        import json
        from smartbids.views.auth_views import enviar_codigo_login

        req = self.rf.post(
            '/api/enviar-codigo-login/',
            data=b'{"email": "admin"}',
            content_type='application/json'
        )
        resp = enviar_codigo_login(req)
        self.assertEqual(resp.status_code, 400)
        data = json.loads(resp.content.decode('utf-8'))
        self.assertEqual(data.get('status'), 'error')
        self.assertIn('formato del correo ingresado no es válido', data.get('mensaje', ''))


class TemplateIngresoTests(SimpleTestCase):
    def test_renderizado_boton_ayuda_y_panel_opciones(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/ingreso/')

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('registro/ingreso.html', request=req)

        # Botón de ayuda en la posición de login actions
        self.assertIn('id="btn-ayuda-login"', html)
        self.assertIn('Ayuda', html)

        # Panel de ayuda con ambas opciones
        self.assertIn('id="ayuda-login-panel"', html)
        self.assertIn('id="forgot-password-link"', html)
        self.assertIn('Recuperar contraseña', html)
        self.assertIn('id="btn-reenviar-activacion"', html)
        self.assertIn('Reenviar activación', html)


class FooterVersionamientoTests(SimpleTestCase):
    def test_footer_info_context_provee_version(self):
        """Verifica que footer_info_context entregue app_version al contexto global."""
        from unittest.mock import patch
        from django.test import RequestFactory
        from smartbids.context_processors import footer_info_context

        rf = RequestFactory()
        req = rf.get('/')
        with patch('smartbids.context_processors.obtener_correos_sistema', return_value={}), \
             patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            ctx = footer_info_context(req)

        self.assertIn('app_version', ctx)
        self.assertEqual(ctx['app_version'], 'v1.0.0')

    def test_footer_template_renderiza_versionamiento_sutil(self):
        """Verifica que el footer incluya la clase site-footer__version y el texto de versión."""
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/')
        with patch('smartbids.context_processors.obtener_correos_sistema', return_value={}), \
             patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('smartbids/includes/footer.html', {'app_version': 'v1.0.0'}, request=req)

        self.assertIn('site-footer__version', html)
        self.assertIn('v1.0.0', html)


class AdministracionGeneralTests(SimpleTestCase):
    def test_navbar_no_contiene_administracion_general(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/mis-licitaciones/')

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('smartbids/Mis_licitaciones.html', request=req, context={'licitaciones': [], 'total_licitaciones': 0})

        self.assertNotIn('id="admin-general-nav"', html)
        self.assertNotIn('Administración General</a>', html)

    def test_perfil_contiene_seccion_admin_en_sesiones(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/perfil/')

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('smartbids/perfil.html', request=req)

        self.assertIn('id="perfil-admin-section"', html)
        self.assertIn('Administración General', html)

    def test_acceso_denegado_no_administrador(self):
        from unittest.mock import patch, MagicMock
        from django.test import RequestFactory
        from smartbids.views.page_views import administracion_general_view

        rf = RequestFactory()
        req = rf.get('/administracion-general/')

        mock_susc = MagicMock()
        mock_susc.codigo_estado = MagicMock(codigo_estado=2, nombre_estado='Habilitado')

        with patch('smartbids.views.page_views.resolver_suscriptor', return_value=mock_susc):
            resp = administracion_general_view(req)
            self.assertEqual(resp.status_code, 302)
            self.assertEqual(resp.url, '/mis-licitaciones/')

    def test_acceso_denegado_anonimo(self):
        from unittest.mock import patch
        from django.test import RequestFactory
        from smartbids.views.page_views import administracion_general_view

        rf = RequestFactory()
        req = rf.get('/administracion-general/')

        with patch('smartbids.views.page_views.resolver_suscriptor', return_value=None):
            resp = administracion_general_view(req)
            self.assertEqual(resp.status_code, 302)
            self.assertEqual(resp.url, '/ingreso/')

    def test_autorizacion_suscriptor_token_valido(self):
        from unittest.mock import MagicMock
        from django.test import RequestFactory
        from smartbids.views.auth_views import validar_autorizacion_suscriptor

        rf = RequestFactory()
        req = rf.post('/api/actualizar-perfil/', HTTP_X_SESSION_TOKEN='token_secreto_123')
        mock_susc = MagicMock(token_sesion='token_secreto_123', firebase_uid='uid_1')

        ok, msg = validar_autorizacion_suscriptor(req, mock_susc)
        self.assertTrue(ok)

    def test_autorizacion_suscriptor_token_invalido(self):
        from unittest.mock import MagicMock
        from django.test import RequestFactory
        from smartbids.views.auth_views import validar_autorizacion_suscriptor

        rf = RequestFactory()
        req = rf.post('/api/actualizar-perfil/', HTTP_X_SESSION_TOKEN='token_atacante')
        mock_susc = MagicMock(token_sesion='token_secreto_123', firebase_uid='uid_1')

        ok, msg = validar_autorizacion_suscriptor(req, mock_susc)
        self.assertFalse(ok)
        self.assertIn('Acceso no autorizado', msg)


class AccionCuentaTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        self.rf = RequestFactory()

    def test_construir_enlace_accion_web(self):
        from smartbids.views.auth_views import _construir_enlace_accion_web
        raw_link = "https://smartbids-auth-users.firebaseapp.com/__/auth/action?mode=verifyEmail&oobCode=XYZ123&apiKey=KEY789&lang=es-419"
        req = self.rf.get('/api/algo/')

        link = _construir_enlace_accion_web(raw_link, request=req, mode_override='verifyEmail')
        self.assertIn('/auth/action/?', link)
        self.assertIn('mode=verifyEmail', link)
        self.assertIn('oobCode=XYZ123', link)
        self.assertIn('apiKey=KEY789', link)

    def test_solicitar_recuperacion_password_validacion(self):
        from smartbids.views.auth_views import solicitar_recuperacion_password
        # Sin email
        req = self.rf.post('/api/solicitar-recuperacion-password/', data=b'{}', content_type='application/json')
        resp = solicitar_recuperacion_password(req)
        self.assertEqual(resp.status_code, 400)

        # Formato inválido
        req_bad = self.rf.post('/api/solicitar-recuperacion-password/', data=b'{"email": "invalido"}', content_type='application/json')
        resp_bad = solicitar_recuperacion_password(req_bad)
        self.assertEqual(resp_bad.status_code, 400)

    def test_solicitar_recuperacion_password_exito(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.auth_views import solicitar_recuperacion_password
        from django.core.cache import cache

        cache.clear()
        mock_user = MagicMock(email='usuario@smartbids.cl')
        with patch('smartbids.views.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', return_value=mock_user), \
             patch('firebase_admin.auth.generate_password_reset_link', return_value='https://firebase.com?oobCode=RESET123&apiKey=KEY'), \
             patch('smartbids.views.auth_views.send_mail') as mock_mail:

            req = self.rf.post(
                '/api/solicitar-recuperacion-password/',
                data=b'{"email": "usuario@smartbids.cl"}',
                content_type='application/json'
            )
            resp = solicitar_recuperacion_password(req)
            self.assertEqual(resp.status_code, 200)
            mock_mail.assert_called_once()
            call_kwargs = mock_mail.call_args.kwargs
            self.assertIn('/auth/action/?', call_kwargs.get('html_message', ''))
            self.assertIn('mode=resetPassword', call_kwargs.get('html_message', ''))

    def test_renderizado_vista_accion_cuenta(self):
        from unittest.mock import patch
        from smartbids.views.page_views import accion_cuenta_view
        req = self.rf.get('/auth/action/?mode=verifyEmail&oobCode=TEST1234')
        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            resp = accion_cuenta_view(req)
        self.assertEqual(resp.status_code, 200)
        content = resp.content.decode('utf-8')
        self.assertIn('btn-activar-cuenta', content)
        self.assertIn('reset-password-form', content)


class ValidacionRegistroEmailTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        self.rf = RequestFactory()

    def test_verificar_existencia_gmail_formato(self):
        from smartbids.views.auth_views import verificar_existencia_gmail
        ok, err = verificar_existencia_gmail('abc@gmail.com')
        self.assertFalse(ok)
        self.assertIn('al menos 6 caracteres', err)

        ok2, err2 = verificar_existencia_gmail('usuarioextremadamentelargoconmasde30caracteres@gmail.com')
        self.assertFalse(ok2)
        self.assertIn('más de 30 caracteres', err2)

        ok3, err3 = verificar_existencia_gmail('contacto@miempresa.cl')
        self.assertTrue(ok3)
        self.assertIsNone(err3)

    def test_validar_email_registro_endpoints(self):
        from unittest.mock import patch, MagicMock
        import firebase_admin.auth
        from smartbids.views.auth_views import validar_email_registro

        req = self.rf.post('/api/validar-email-registro/', data=b'{}', content_type='application/json')
        resp = validar_email_registro(req)
        self.assertEqual(resp.status_code, 400)

        with patch('smartbids.views.auth_views.verificar_existencia_gmail', return_value=(False, 'La cuenta no existe en Gmail')):
            req_gmail = self.rf.post('/api/validar-email-registro/', data=b'{"email": "inexistente@gmail.com"}', content_type='application/json')
            resp_gmail = validar_email_registro(req_gmail)
            self.assertEqual(resp_gmail.status_code, 400)
            self.assertIn('no existe en Gmail', resp_gmail.content.decode('utf-8'))

        mock_user = MagicMock(email='existente@smartbids.cl')
        with patch('smartbids.views.auth_views.verificar_existencia_gmail', return_value=(True, None)), \
             patch('smartbids.views.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', return_value=mock_user):
            req_duplicado = self.rf.post('/api/validar-email-registro/', data=b'{"email": "existente@smartbids.cl"}', content_type='application/json')
            resp_duplicado = validar_email_registro(req_duplicado)
            self.assertEqual(resp_duplicado.status_code, 400)
            self.assertIn('Ya existe una cuenta registrada', resp_duplicado.content.decode('utf-8'))

        with patch('smartbids.views.auth_views.verificar_existencia_gmail', return_value=(True, None)), \
             patch('smartbids.views.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', side_effect=firebase_admin.auth.UserNotFoundError('No user')):
            req_nuevo = self.rf.post('/api/validar-email-registro/', data=b'{"email": "nuevo@smartbids.cl"}', content_type='application/json')
            resp_nuevo = validar_email_registro(req_nuevo)
            self.assertEqual(resp_nuevo.status_code, 200)

    def test_filtrar_licitaciones_sin_perfil_match_cero(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs

        with patch('smartbids.views.licitaciones_views.Licitacion.objects', mock_qs), \
             patch('smartbids.views.licitaciones_views.Preferencia.objects.filter') as mock_pref:
            # Suscriptor sin preferencias configuradas (None o listas vacías)
            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=[],
                pref_tipo_licitacion=[],
                pref_ucom=[],
                pref_palabras_claves=[],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/')
            qs, match_perfil = filtrar_licitaciones(req, suscriptor=MagicMock())

            # El match debe ser 0 y mock_qs no debe haberse filtrado por q_perfil
            self.assertEqual(match_perfil, 0)
            mock_qs.filter.assert_not_called()

    def test_filtrar_licitaciones_con_perfil_aplica_filtro_y_aumenta_match(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_filtered_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs
        mock_qs.filter.return_value.distinct.return_value = mock_filtered_qs

        with patch('smartbids.views.licitaciones_views.Licitacion.objects') as mock_licitacion_objects, \
             patch('smartbids.views.licitaciones_views.Preferencia.objects.filter') as mock_pref:
            mock_licitacion_objects.select_related.return_value.all.return_value.order_by.return_value = mock_qs
            mock_licitacion_objects.filter.return_value.distinct.return_value.count.return_value = 42

            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=['13101'],
                pref_tipo_licitacion=['LE'],
                pref_ucom=[],
                pref_palabras_claves=['servidor', 'computador'],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/?tab=claves')
            qs, match_perfil = filtrar_licitaciones(req, suscriptor=MagicMock())

            # El match con perfil debe ser 42 y el queryset devuelto debe estar filtrado en pestaña de claves
            self.assertEqual(match_perfil, 42)
            self.assertEqual(qs, mock_filtered_qs)

    def test_filtrar_licitaciones_tab_todas_sin_filtro_de_claves(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs

        with patch('smartbids.views.licitaciones_views.Licitacion.objects') as mock_licitacion_objects, \
             patch('smartbids.views.licitaciones_views.Preferencia.objects.filter') as mock_pref:
            mock_licitacion_objects.select_related.return_value.all.return_value.order_by.return_value = mock_qs
            mock_licitacion_objects.filter.return_value.distinct.return_value.count.return_value = 15

            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=[],
                pref_tipo_licitacion=[],
                pref_ucom=[],
                pref_palabras_claves=['software'],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/?tab=todas')
            qs, match_perfil = filtrar_licitaciones(req, suscriptor=MagicMock())

            # En pestaña Todas, no se aplica el filtro de palabras clave a qs, pero se calcula el match
            self.assertEqual(match_perfil, 15)
            self.assertEqual(qs, mock_qs)

    def test_mis_licitaciones_view_kpi_match_cero_cuando_perfil_vacio(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.licitaciones_views import mis_licitaciones_view

        with patch('smartbids.views.licitaciones_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.views.licitaciones_views.cache.get_or_set', return_value=62650), \
             patch('smartbids.views.licitaciones_views.filtrar_licitaciones', return_value=(MagicMock(), 0)), \
             patch('smartbids.views.licitaciones_views.render') as mock_render, \
             patch('smartbids.views.licitaciones_views.Paginator') as mock_paginator:
            
            mock_paginator.return_value.get_page.return_value = []
            req = self.rf.get('/mis-licitaciones/')
            mis_licitaciones_view(req)

            self.assertTrue(mock_render.called)
            context = mock_render.call_args[0][2]
            self.assertEqual(context['match_perfil'], 0)


class ResumenDiaTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def test_resumen_dia_view_sin_suscriptor(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.resumen_views import resumen_dia_view

        mock_lic_qs = MagicMock()
        mock_lic_qs.count.return_value = 0
        mock_lic_qs.__iter__.return_value = []
        mock_lic_qs.__getitem__.return_value = []

        with patch('smartbids.views.resumen_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.views.resumen_views.render') as mock_render, \
             patch('smartbids.views.resumen_views.Licitacion.objects') as mock_lic:
            mock_lic.filter.return_value.select_related.return_value.order_by.return_value = mock_lic_qs
            req = self.rf.get('/resumen/')
            resumen_dia_view(req)

            self.assertTrue(mock_render.called)
            template_name = mock_render.call_args[0][1]
            context = mock_render.call_args[0][2]
            self.assertEqual(template_name, 'smartbids/resumen_dia.html')
            self.assertEqual(context['nombre_usuario'], 'Suscriptor')
            self.assertEqual(context['total_coincidencias'], 0)
            self.assertEqual(context['total_seguidas'], 0)

    def test_resumen_dia_view_con_suscriptor(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.resumen_views import resumen_dia_view

        mock_susc = MagicMock(sus_nombre1='Juan', sus_apellido1='Pérez')
        mock_pref = MagicMock(pref_palabras_claves=['servicios'], pref_productos=[])

        mock_lic_qs = MagicMock()
        mock_lic_qs.count.return_value = 3
        mock_lic_qs.__iter__.return_value = []
        mock_lic_qs.__getitem__.return_value = []

        mock_destacadas_qs = MagicMock()
        mock_destacadas_qs.__iter__.return_value = []
        mock_destacadas_qs.__getitem__.return_value = []

        mock_notif_qs = MagicMock()
        mock_notif_qs.filter.return_value.count.return_value = 1
        mock_notif_qs.__iter__.return_value = []
        mock_notif_qs.__getitem__.return_value = []

        with patch('smartbids.views.resumen_views.resolver_suscriptor', return_value=mock_susc), \
             patch('smartbids.views.resumen_views._sincronizar_seguimientos'), \
             patch('smartbids.views.resumen_views.Preferencia.objects.filter') as mock_pref_filter, \
             patch('smartbids.views.resumen_views.SeguimientoLicitacion.objects.filter') as mock_seg_filter, \
             patch('smartbids.views.resumen_views.NotificacionLicitacion.objects.filter') as mock_notif_filter, \
             patch('smartbids.views.resumen_views.Licitacion.objects') as mock_lic, \
             patch('smartbids.views.resumen_views.render') as mock_render:
            
            mock_pref_filter.return_value.first.return_value = mock_pref
            mock_seg_filter.return_value.select_related.return_value.count.return_value = 2
            mock_seg_filter.return_value.select_related.return_value.values_list.return_value = ['LIC-1']
            mock_seg_filter.return_value.select_related.return_value.__iter__.return_value = []
            mock_seg_filter.return_value.select_related.return_value.__getitem__.return_value = []
            mock_notif_filter.return_value.select_related.return_value.order_by.return_value = mock_notif_qs
            mock_lic.filter.return_value.distinct.return_value.count.return_value = 8
            mock_lic.filter.return_value.select_related.return_value.order_by.return_value = mock_lic_qs
            mock_lic.filter.return_value.filter.return_value.select_related.return_value.order_by.return_value.distinct.return_value = mock_destacadas_qs

            req = self.rf.get('/resumen/')
            resumen_dia_view(req)

            self.assertTrue(mock_render.called)
            context = mock_render.call_args[0][2]
            self.assertEqual(context['nombre_usuario'], 'Juan Pérez')
            self.assertTrue(context['tiene_palabras_perfil'])
            self.assertEqual(context['total_coincidencias'], 8)

    def test_resumen_dia_view_oportunidades_y_cierran_pronto_solo_match_perfil(self):
        """Verifica que Oportunidades del Día y Cierran Pronto filtren exclusivamente por el perfil configurado."""
        from unittest.mock import patch, MagicMock
        from smartbids.views.resumen_views import resumen_dia_view

        mock_susc = MagicMock(sus_nombre1='Ana')
        mock_pref = MagicMock(pref_tipo_licitacion=['LE'], pref_palabras_claves=[], pref_productos=[], pref_comunas=[], pref_ucom=[])

        mock_match_lic = MagicMock(lic_codigo='MATCH-1', lic_tipo_licitacion='LE')
        mock_match_qs = MagicMock()
        mock_match_qs.__iter__.return_value = [mock_match_lic]
        mock_match_qs.__getitem__.return_value = [mock_match_lic]
        mock_match_qs.count.return_value = 1

        with patch('smartbids.views.resumen_views.resolver_suscriptor', return_value=mock_susc), \
             patch('smartbids.views.resumen_views._sincronizar_seguimientos'), \
             patch('smartbids.views.resumen_views.Preferencia.objects.filter') as mock_pref_filter, \
             patch('smartbids.views.resumen_views.SeguimientoLicitacion.objects.filter') as mock_seg_filter, \
             patch('smartbids.views.resumen_views.NotificacionLicitacion.objects.filter') as mock_notif_filter, \
             patch('smartbids.views.resumen_views.Licitacion.objects') as mock_lic, \
             patch('smartbids.views.resumen_views.estado_licitacion', return_value=('ofertas', 'Recepción de ofertas')), \
             patch('smartbids.views.resumen_views.render') as mock_render:

            mock_pref_filter.return_value.first.return_value = mock_pref
            mock_seg_filter.return_value.select_related.return_value.count.return_value = 0
            mock_seg_filter.return_value.select_related.return_value.values_list.return_value = []
            mock_seg_filter.return_value.select_related.return_value.__iter__.return_value = []
            mock_seg_filter.return_value.select_related.return_value.__getitem__.return_value = []
            mock_notif_filter.return_value.select_related.return_value.order_by.return_value.filter.return_value.count.return_value = 0
            mock_notif_filter.return_value.select_related.return_value.order_by.return_value.__iter__.return_value = []

            mock_lic.filter.return_value.distinct.return_value.count.return_value = 1
            mock_lic.filter.return_value.select_related.return_value.order_by.return_value.distinct.return_value = mock_match_qs
            mock_lic.filter.return_value.filter.return_value.select_related.return_value.order_by.return_value.distinct.return_value = mock_match_qs

            req = self.rf.get('/resumen/')
            resumen_dia_view(req)

            self.assertTrue(mock_render.called)
            context = mock_render.call_args[0][2]
            self.assertEqual(len(context['licitaciones_destacadas']), 1)
            self.assertEqual(context['licitaciones_destacadas'][0].lic_codigo, 'MATCH-1')
            self.assertEqual(len(context['licitaciones_urgentes']), 1)
            self.assertEqual(context['licitaciones_urgentes'][0].lic_codigo, 'MATCH-1')

    def test_resumen_dia_view_cierran_pronto_fallback_mercado(self):
        """Verifica que si no hay cierres en 72h del perfil, Cierran Pronto use las oportunidades activas del mercado."""
        from unittest.mock import patch, MagicMock
        from smartbids.views.resumen_views import resumen_dia_view

        mock_susc = MagicMock(sus_nombre1='Carlos')
        mock_pref = MagicMock(pref_tipo_licitacion=['LE'], pref_palabras_claves=[], pref_productos=[], pref_comunas=[], pref_ucom=[])

        mock_general_lic = MagicMock(lic_codigo='GEN-1', lic_tipo_licitacion='LE')
        mock_general_qs = MagicMock()
        mock_general_qs.__iter__.return_value = [mock_general_lic]
        mock_general_qs.__getitem__.return_value = [mock_general_lic]
        mock_general_qs.count.return_value = 3

        mock_empty_qs = MagicMock()
        mock_empty_qs.count.return_value = 0
        mock_empty_qs.__iter__.return_value = []
        mock_empty_qs.__getitem__.return_value = []

        with patch('smartbids.views.resumen_views.resolver_suscriptor', return_value=mock_susc), \
             patch('smartbids.views.resumen_views._sincronizar_seguimientos'), \
             patch('smartbids.views.resumen_views.Preferencia.objects.filter') as mock_pref_filter, \
             patch('smartbids.views.resumen_views.SeguimientoLicitacion.objects.filter') as mock_seg_filter, \
             patch('smartbids.views.resumen_views.NotificacionLicitacion.objects.filter') as mock_notif_filter, \
             patch('smartbids.views.resumen_views.Licitacion.objects') as mock_lic, \
             patch('smartbids.views.resumen_views.estado_licitacion', return_value=('ofertas', 'Recepción de ofertas')), \
             patch('smartbids.views.resumen_views.render') as mock_render:

            mock_pref_filter.return_value.first.return_value = mock_pref
            mock_seg_filter.return_value.select_related.return_value.count.return_value = 0
            mock_seg_filter.return_value.select_related.return_value.values_list.return_value = []
            mock_seg_filter.return_value.select_related.return_value.__iter__.return_value = []
            mock_seg_filter.return_value.select_related.return_value.__getitem__.return_value = []
            mock_notif_filter.return_value.select_related.return_value.order_by.return_value.filter.return_value.count.return_value = 0
            mock_notif_filter.return_value.select_related.return_value.order_by.return_value.__iter__.return_value = []

            mock_lic.filter.return_value.distinct.return_value.count.return_value = 15
            mock_lic.filter.return_value.filter.return_value.select_related.return_value.order_by.return_value.distinct.return_value = mock_empty_qs
            mock_lic.filter.return_value.select_related.return_value.order_by.return_value.distinct.return_value = mock_general_qs

            req = self.rf.get('/resumen/')
            resumen_dia_view(req)

            self.assertTrue(mock_render.called)
            context = mock_render.call_args[0][2]
            self.assertTrue(context['cierre_es_general'])
            self.assertEqual(context['total_cierre_proximo'], 3)


class AdminSafetyAndEmailConfigTests(SimpleTestCase):
    def setUp(self):
        import json
        from django.test import RequestFactory
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def test_admin_auto_eliminacion_bloqueada(self):
        import json
        from unittest.mock import patch, MagicMock
        from smartbids.views.firebase_policy_views import usuarios_firebase

        mock_susc = MagicMock(sus_nombre1='Admin', sus_apellido1='Uno', sus_nombre_social='')

        with patch('smartbids.views.firebase_policy_views._usuario_es_administrador', return_value=True), \
             patch('smartbids.views.firebase_policy_views._firebase_admin_session'), \
             patch('smartbids.views.firebase_policy_views._firebase_accounts_url'), \
             patch('smartbids.views.firebase_policy_views.Suscriptor.objects') as mock_susc_objs:
            
            mock_susc_objs.select_related.return_value.filter.return_value.first.return_value = mock_susc
            req = self.rf.delete('/api/firebase/usuarios/?uid=admin_123', HTTP_X_FIREBASE_UID='admin_123')
            resp = usuarios_firebase(req)
            self.assertEqual(resp.status_code, 403)
            data = json.loads(resp.content)
            self.assertIn('no puede eliminar su propia cuenta', data['mensaje'])

    def test_admin_auto_remocion_bloqueada(self):
        import json
        from unittest.mock import patch, MagicMock
        from smartbids.views.firebase_policy_views import usuarios_firebase

        mock_susc = MagicMock(sus_nombre1='Admin', sus_apellido1='Uno', sus_nombre_social='')
        mock_estado_suscriptor = MagicMock(nombre_estado='suscriptor')

        with patch('smartbids.views.firebase_policy_views._usuario_es_administrador', return_value=True), \
             patch('smartbids.views.firebase_policy_views._firebase_admin_session'), \
             patch('smartbids.views.firebase_policy_views._firebase_accounts_url'), \
             patch('smartbids.views.firebase_policy_views.Suscriptor.objects') as mock_susc_objs, \
             patch('smartbids.views.firebase_policy_views.EstadoSuscriptor.objects.filter', return_value=MagicMock(first=lambda: mock_estado_suscriptor)):
            
            mock_susc_objs.select_related.return_value.filter.return_value.first.return_value = mock_susc
            body = json.dumps({'uid': 'admin_123', 'codigo_estado': 2})
            req = self.rf.patch('/api/firebase/usuarios/', data=body, content_type='application/json', HTTP_X_FIREBASE_UID='admin_123')
            resp = usuarios_firebase(req)
            self.assertEqual(resp.status_code, 403)
            data = json.loads(resp.content)
            self.assertIn('no puede revocar su propio rol', data['mensaje'])

    def test_admin_eliminar_otro_admin_genera_auditoria_y_correo(self):
        from unittest.mock import patch, MagicMock
        from smartbids.views.firebase_policy_views import usuarios_firebase

        mock_target_susc = MagicMock(
            codigo_estado=MagicMock(nombre_estado='administrador', codigo_estado=5),
            sus_nombre1='Admin',
            sus_apellido1='Dos',
            sus_nombre_social='Admin 2'
        )
        mock_session = MagicMock()
        mock_session.post.return_value = MagicMock(ok=True, status_code=200)

        with patch('smartbids.views.firebase_policy_views._usuario_es_administrador', return_value=True), \
             patch('smartbids.views.firebase_policy_views._firebase_admin_session', return_value=mock_session), \
             patch('smartbids.views.firebase_policy_views._firebase_accounts_url', return_value='https://identitytoolkit.googleapis.com'), \
             patch('smartbids.views.firebase_policy_views.Suscriptor.objects') as mock_susc_objs, \
             patch('smartbids.views.firebase_policy_views.Preferencia.objects'), \
             patch('smartbids.views.firebase_policy_views.registrar_auditoria_admin') as mock_registrar, \
             patch('smartbids.views.firebase_policy_views.notificar_accion_entre_administradores') as mock_notificar:
            
            mock_susc_objs.select_related.return_value.filter.return_value.first.return_value = mock_target_susc
            mock_susc_objs.filter.return_value.first.return_value = mock_target_susc

            req = self.rf.delete('/api/firebase/usuarios/?uid=admin_456&email=admin2@empresa.cl', HTTP_X_FIREBASE_UID='admin_123')
            resp = usuarios_firebase(req)
            self.assertEqual(resp.status_code, 200)
            self.assertTrue(mock_registrar.called)
            self.assertEqual(mock_registrar.call_args[1]['accion'], 'ELIMINACION')
            self.assertTrue(mock_notificar.called)
            self.assertEqual(mock_notificar.call_args[1]['afectado_email'], 'admin2@empresa.cl')

    def test_config_correo_get_y_patch(self):
        import json
        from unittest.mock import patch, MagicMock
        from smartbids.views.mensajeria_views import administrar_config_correo

        with patch('smartbids.views.mensajeria_views._usuario_es_administrador', return_value=True), \
             patch('smartbids.views.mensajeria_views.ConfiguracionCorreo.objects') as mock_cfg_objs:
            
            mock_cfg = MagicMock(
                email_contacto='contacto@chileavanza.cl',
                email_notificaciones='notif@chileavanza.cl',
                nombre_remitente='SmartBids',
                email_seguridad_auditoria='auditoria@chileavanza.cl'
            )
            mock_cfg_objs.first.return_value = mock_cfg

            # Test GET
            req_get = self.rf.get('/api/administracion/config-correo/')
            resp_get = administrar_config_correo(req_get)
            self.assertEqual(resp_get.status_code, 200)
            data_get = json.loads(resp_get.content)
            self.assertEqual(data_get['status'], 'ok')

            # Test PATCH
            body = json.dumps({
                'email_contacto': 'nuevo_contacto@smartbids.cl',
                'email_notificaciones': 'alertas@smartbids.cl',
                'nombre_remitente': 'SmartBids Admin',
                'email_seguridad_auditoria': 'seguridad@smartbids.cl'
            })
            req_patch = self.rf.patch('/api/administracion/config-correo/', data=body, content_type='application/json')
            resp_patch = administrar_config_correo(req_patch)
            self.assertEqual(resp_patch.status_code, 200)
            data_patch = json.loads(resp_patch.content)
            self.assertEqual(data_patch['config']['email_contacto'], 'nuevo_contacto@smartbids.cl')

    def test_auditoria_admin_view_retorna_registros(self):
        import json
        from django.utils import timezone
        from unittest.mock import patch, MagicMock
        from smartbids.views.mensajeria_views import auditoria_admin_view

        mock_reg = MagicMock(
            id=1,
            ejecutor_email='admin1@chileavanza.cl',
            ejecutor_nombre='Admin Uno',
            afectado_email='admin2@chileavanza.cl',
            afectado_nombre='Admin Dos',
            accion='ELIMINACION',
            detalle='Eliminó admin 2',
            creado_el=timezone.now(),
            ip_origen='127.0.0.1'
        )

        with patch('smartbids.views.mensajeria_views._usuario_es_administrador', return_value=True), \
             patch('smartbids.views.mensajeria_views.AuditoriaAdmin.objects.all') as mock_all:
            mock_all.return_value.order_by.return_value = [mock_reg]
            req = self.rf.get('/api/administracion/auditoria-admin/')
            resp = auditoria_admin_view(req)
            self.assertEqual(resp.status_code, 200)
            data = json.loads(resp.content)
            self.assertEqual(len(data['registros']), 1)
            self.assertEqual(data['registros'][0]['accion'], 'ELIMINACION')

    def test_producto_codigo_formateado_standard_y_obras_civiles(self):
        """Verifica que el código se formatee con guiones y soporte primer nivel de 3 dígitos (Obras Civiles)."""
        from smartbids.models.catalog import Producto

        p_standard = Producto(
            codigo_producto='50193001',
            descripcion='Comida para bebés',
            nivel1_id='50',
            nivel2='19',
            nivel3='30',
            nivel4='01'
        )
        self.assertEqual(p_standard.codigo_formateado, '50-19-30-01')
        self.assertIn('50-19-30-01', str(p_standard))

        p_obras = Producto(
            codigo_producto='101101131',
            descripcion='1era',
            nivel1_id='101',
            nivel2='10',
            nivel3='11',
            nivel4='31',
            glosa_nivel2='Obras Mayores',
            glosa_nivel3='Obra Civil',
            glosa_nivel4='1era'
        )
        self.assertEqual(p_obras.codigo_formateado, '101-10-11-31')
        self.assertIn('101-10-11-31', str(p_obras))

    def test_catalogos_preferencias_productos_arbol_jerarquia(self):
        """Verifica la jerarquía Grupo -> Nivel 1 -> Productos y la ausencia de textos erróneos."""
        import json
        from unittest.mock import patch, MagicMock
        from smartbids.views.perfil_views import catalogos_preferencias
        from smartbids.models.catalog import GrupoProducto, Nivel1Producto

        mock_grupos = [
            GrupoProducto(codigo_grupo_producto='01', nombre_grupo_producto='Alimentos'),
            GrupoProducto(codigo_grupo_producto='12', nombre_grupo_producto='Obras Públicas')
        ]
        mock_niveles = [
            Nivel1Producto(codigo_nivel1='50', descripcion='Alimentos, bebidas y tabaco', codigo_grupo_producto_id='01'),
            Nivel1Producto(codigo_nivel1='101', descripcion='Obras', codigo_grupo_producto_id='12')
        ]
        mock_prods = [
            {
                'codigo_producto': '50193001',
                'descripcion': 'Comida para bebés',
                'nivel1_id': '50',
                'nivel2': '19',
                'nivel3': '30',
                'nivel4': '01',
                'glosa_nivel2': 'Alimentos preparados',
                'glosa_nivel3': 'Bebidas infantiles',
                'glosa_nivel4': 'Comida para bebés'
            },
            {
                'codigo_producto': '101101131',
                'descripcion': '1era',
                'nivel1_id': '101',
                'nivel2': '10',
                'nivel3': '11',
                'nivel4': '31',
                'glosa_nivel2': 'Obras Mayores',
                'glosa_nivel3': 'Obra Civil',
                'glosa_nivel4': '1era'
            }
        ]

        with patch('smartbids.views.perfil_views.GrupoProducto.objects.all') as m_g_all, \
             patch('smartbids.views.perfil_views.Nivel1Producto.objects.all') as m_n_all, \
             patch('smartbids.views.perfil_views.Producto.objects.exclude') as m_p_exc:

            m_g_all.return_value.order_by.return_value = mock_grupos
            m_n_all.return_value.order_by.return_value = mock_niveles

            m_p_exc.return_value.exclude.return_value.order_by.return_value.values.return_value = mock_prods

            req = self.rf.get('/api/catalogos-preferencias/?tipo=productos_arbol')
            resp = catalogos_preferencias(req)
            self.assertEqual(resp.status_code, 200)

            data = json.loads(resp.content)
            self.assertEqual(data.get('status'), 'ok')
            self.assertIn('grupos', data)
            self.assertIn('rubros', data)

            grupos = data['grupos']
            self.assertEqual(len(grupos), 2)

            nombres_grupos = [g['nombre_grupo'] for g in grupos]
            self.assertIn('Alimentos', nombres_grupos)
            self.assertIn('Obras Públicas', nombres_grupos)

            # Verificar Obras Civiles con código de 9 dígitos y primer nivel de 3 dígitos (101)
            grupo_obras = next((g for g in grupos if g['nombre_grupo'] == 'Obras Públicas'), None)
            self.assertIsNotNone(grupo_obras)
            self.assertEqual(grupo_obras['niveles'][0]['codigo_nivel1'], '101')
            prod_obras = grupo_obras['niveles'][0]['productos'][0]
            self.assertEqual(prod_obras['codigo_producto'], '101101131')
            self.assertEqual(prod_obras['codigo_formateado'], '101-10-11-31')
            self.assertEqual(prod_obras['descripcion'], '1era')

            # Verificar estándar ONU 8 dígitos
            grupo_alimentos = next((g for g in grupos if g['nombre_grupo'] == 'Alimentos'), None)
            self.assertIsNotNone(grupo_alimentos)
            prod_alim = grupo_alimentos['niveles'][0]['productos'][0]
            self.assertEqual(prod_alim['codigo_producto'], '50193001')
            self.assertEqual(prod_alim['codigo_formateado'], '50-19-30-01')

            # Verificar que no existan textos erróneos
            texto_completo = json.dumps(data)
            self.assertNotIn('Marvel', texto_completo)
            self.assertNotIn('Segmento ONU Marvel', texto_completo)
            self.assertNotIn('Segmento ONU 48', texto_completo)
            self.assertNotIn('Segmento ONU 54', texto_completo)

    def test_catalogos_preferencias_busqueda_productos(self):
        """Verifica que el filtrado por q busque en descripciones y códigos de productos."""
        import json
        from unittest.mock import patch
        from smartbids.views.perfil_views import catalogos_preferencias
        from smartbids.models.catalog import GrupoProducto, Nivel1Producto

        mock_grupos = [GrupoProducto(codigo_grupo_producto='01', nombre_grupo_producto='Alimentos')]
        mock_niveles = [Nivel1Producto(codigo_nivel1='50', descripcion='Alimentos, bebidas y tabaco', codigo_grupo_producto_id='01')]
        mock_prods = [{
            'codigo_producto': '50193001',
            'descripcion': 'Comida para bebés',
            'nivel1_id': '50',
            'nivel2': '19',
            'nivel3': '30',
            'nivel4': '01',
            'glosa_nivel2': '',
            'glosa_nivel3': '',
            'glosa_nivel4': ''
        }]

        with patch('smartbids.views.perfil_views.GrupoProducto.objects.all') as m_g_all, \
             patch('smartbids.views.perfil_views.Nivel1Producto.objects.all') as m_n_all, \
             patch('smartbids.views.perfil_views.Producto.objects.exclude') as m_p_exc:

            m_g_all.return_value.order_by.return_value = mock_grupos
            m_n_all.return_value.order_by.return_value = mock_niveles
            m_p_exc.return_value.exclude.return_value.filter.return_value.__getitem__.return_value.values.return_value = mock_prods

            req = self.rf.get('/api/catalogos-preferencias/?tipo=productos_arbol&q=beb%C3%A9s')
            resp = catalogos_preferencias(req)
            self.assertEqual(resp.status_code, 200)

            data = json.loads(resp.content)
            self.assertEqual(data.get('status'), 'ok')
            self.assertIn('grupos', data)
            self.assertEqual(len(data['grupos']), 1)
            self.assertEqual(data['grupos'][0]['niveles'][0]['productos'][0]['codigo_producto'], '50193001')

    def test_obtener_perfil_productos_detalle_formato(self):
        """Verifica que productos_detalle entregue etiquetas con código formateado y descripción adecuada."""
        import json
        from unittest.mock import patch, MagicMock
        from smartbids.views.perfil_views import obtener_perfil_suscriptor

        mock_susc = MagicMock(
            id_suscriptor=10,
            firebase_uid='uid_test_10',
            sus_nombre1='Test',
            sus_nombre2='',
            sus_apellido1='User',
            sus_apellido2='',
            sus_nombre_social='',
            sus_iniciales='TU',
            sus_rut_empresa=None,
            token_sesion='',
            codigo_estado=None,
            codigo_estado_id=None,
            fecha_registro=None,
            fecha_actualizacion=None
        )
        mock_pref = MagicMock(
            pref_comunas=[],
            pref_tipo_licitacion=[],
            pref_ucom=[],
            pref_palabras_claves=[],
            pref_productos=['50193001', '101101131']
        )
        mock_prods = [
            {
                'codigo_producto': '50193001',
                'descripcion': 'Comida para bebés',
                'nivel1_id': '50',
                'nivel2': '19',
                'nivel3': '30',
                'nivel4': '01',
                'glosa_nivel3': '',
                'glosa_nivel4': ''
            },
            {
                'codigo_producto': '101101131',
                'descripcion': '1era',
                'nivel1_id': '101',
                'nivel2': '10',
                'nivel3': '11',
                'nivel4': '31',
                'glosa_nivel3': 'Obra Civil',
                'glosa_nivel4': '1era'
            }
        ]

        m_susc_qs = MagicMock()
        m_susc_qs.select_related.return_value.filter.return_value.first.return_value = mock_susc
        m_pref_qs = MagicMock()
        m_pref_qs.filter.return_value.first.return_value = mock_pref
        m_com_qs = MagicMock()
        m_com_qs.filter.return_value.values.return_value = []
        m_ucom_qs = MagicMock()
        m_ucom_qs.filter.return_value.values.return_value = []
        m_prod_qs = MagicMock()
        m_prod_qs.filter.return_value.values.return_value = mock_prods

        with patch('smartbids.views.perfil_views.Suscriptor.objects', m_susc_qs), \
             patch('smartbids.views.perfil_views.Preferencia.objects', m_pref_qs), \
             patch('smartbids.views.perfil_views.Comuna.objects', m_com_qs), \
             patch('smartbids.views.perfil_views.UnidadCompra.objects', m_ucom_qs), \
             patch('smartbids.views.perfil_views.Producto.objects', m_prod_qs):

            req = self.rf.post('/api/perfil/obtener/', data=b'{"uid": "uid_test_10"}', content_type='application/json')
            resp = obtener_perfil_suscriptor(req)
            self.assertEqual(resp.status_code, 200)

            data = json.loads(resp.content)
            self.assertEqual(data.get('status'), 'ok')
            prods_det = data['datos']['preferencias']['productos_detalle']
            self.assertEqual(len(prods_det), 2)

            self.assertEqual(prods_det[0]['code'], '50193001')
            self.assertEqual(prods_det[0]['label'], 'Comida para bebés (50-19-30-01)')

            self.assertEqual(prods_det[1]['code'], '101101131')
            self.assertEqual(prods_det[1]['label'], '1era (101-10-11-31)')


class TipoProcedimientoLicitacionTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def tearDown(self):
        from django.core.cache import cache
        cache.clear()

    def test_obtener_catalogo_tipos_licitacion_estructura_y_clasificacion(self):
        """Verifica que el catálogo organice 5 tramos públicos (L1, LE, LP, LQ, LR) y 3 privados (E2, CO, B2)."""
        from unittest.mock import patch, MagicMock
        from smartbids.utils.licitacion_utils import obtener_catalogo_tipos_licitacion

        mock_cat_pub = MagicMock(codigo_cat_licitacion=1, nombre_cat_licitacion='Pública')
        mock_cat_priv = MagicMock(codigo_cat_licitacion=2, nombre_cat_licitacion='Privada')
        mock_tramo_100 = MagicMock(codigo_tramo_licitacion=2, nombre_tramo_licitacion='menor a 100 UTM')
        mock_tramo_1000 = MagicMock(codigo_tramo_licitacion=3, nombre_tramo_licitacion='100 a 1000 UTM')
        mock_tramo_2000 = MagicMock(codigo_tramo_licitacion=4, nombre_tramo_licitacion='1000 a 2000 UTM')
        mock_tramo_5000_tramo = MagicMock(codigo_tramo_licitacion=5, nombre_tramo_licitacion='2000 a 5000 UTM')
        mock_tramo_5000 = MagicMock(codigo_tramo_licitacion=6, nombre_tramo_licitacion='mayor a 5000 UTM')

        mock_tl_list = [
            MagicMock(codigo_tipo_licitacion='L1', glosa='Licitación pública inferior a 100 UTM', codigo_cat_licitacion=mock_cat_pub, codigo_cat_licitacion_id=1, codigo_tramo_licitacion=mock_tramo_100),
            MagicMock(codigo_tipo_licitacion='LE', glosa='Licitación Pública igual o superior a 100 UTM e inferior a 1.000 UTM', codigo_cat_licitacion=mock_cat_pub, codigo_cat_licitacion_id=1, codigo_tramo_licitacion=mock_tramo_1000),
            MagicMock(codigo_tipo_licitacion='LP', glosa='Licitación Pública igual o superior a 1.000 UTM e inferior a 2.000 UTM', codigo_cat_licitacion=mock_cat_pub, codigo_cat_licitacion_id=1, codigo_tramo_licitacion=mock_tramo_2000),
            MagicMock(codigo_tipo_licitacion='LQ', glosa='Licitación Pública igual o superior a 2.000 UTM e inferior a 5.000 UTM', codigo_cat_licitacion=mock_cat_pub, codigo_cat_licitacion_id=1, codigo_tramo_licitacion=mock_tramo_5000_tramo),
            MagicMock(codigo_tipo_licitacion='LR', glosa='Licitación Pública Mayor a 5000 UTM', codigo_cat_licitacion=mock_cat_pub, codigo_cat_licitacion_id=1, codigo_tramo_licitacion=mock_tramo_5000),
            MagicMock(codigo_tipo_licitacion='E2', glosa='Licitación Privada Inferior a 100 UTM', codigo_cat_licitacion=mock_cat_priv, codigo_cat_licitacion_id=2, codigo_tramo_licitacion=mock_tramo_100),
            MagicMock(codigo_tipo_licitacion='CO', glosa='Licitación Privada igual o superior a 100 UTM e inferior a 1000 UTM', codigo_cat_licitacion=mock_cat_priv, codigo_cat_licitacion_id=2, codigo_tramo_licitacion=mock_tramo_1000),
            MagicMock(codigo_tipo_licitacion='B2', glosa='Licitación Privada igual o superior a 1000 UTM e inferior a 2000 UTM', codigo_cat_licitacion=mock_cat_priv, codigo_cat_licitacion_id=2, codigo_tramo_licitacion=mock_tramo_2000),
        ]

        with patch('smartbids.models.catalog.TipoLicitacion.objects') as m_tl_objects:
            m_tl_objects.select_related.return_value.filter.return_value = mock_tl_list
            catalogo = obtener_catalogo_tipos_licitacion()

        self.assertIn('publicas', catalogo)
        self.assertIn('privadas', catalogo)
        self.assertIn('todos_codigos', catalogo)
        self.assertEqual(catalogo['categorias_publicas_ids'], [1, 6])
        self.assertEqual(catalogo['categorias_privadas_ids'], [2, 7])

        # 5 públicas: L1, LE, LP, LQ, LR
        codigos_pub = [item['codigo'] for item in catalogo['publicas']]
        self.assertEqual(codigos_pub, ['L1', 'LE', 'LP', 'LQ', 'LR'])
        for item in catalogo['publicas']:
            self.assertEqual(item['categoria'], 'Pública')

        # Verificar tramos específicos de LP y LQ
        item_lp = next(item for item in catalogo['publicas'] if item['codigo'] == 'LP')
        self.assertEqual(item_lp['tramo'], '1.000 a 2.000 UTM')
        self.assertIn('1.000 UTM e inferior a 2.000 UTM', item_lp['glosa'])

        item_lq = next(item for item in catalogo['publicas'] if item['codigo'] == 'LQ')
        self.assertEqual(item_lq['tramo'], '2.000 a 5.000 UTM')
        self.assertIn('2.000 UTM e inferior a 5.000 UTM', item_lq['glosa'])

        # 3 privadas: E2, CO, B2
        codigos_priv = [item['codigo'] for item in catalogo['privadas']]
        self.assertEqual(codigos_priv, ['E2', 'CO', 'B2'])
        for item in catalogo['privadas']:
            self.assertEqual(item['categoria'], 'Privada')

        # B2 debe ser Licitación Privada, nunca pública ni 'Compra Ágil'
        item_b2 = next(item for item in catalogo['privadas'] if item['codigo'] == 'B2')
        self.assertEqual(item_b2['categoria'], 'Privada')
        self.assertIn('1.000 UTM', item_b2['tramo'])
        self.assertIn('Licitación Privada', item_b2['glosa'])

        # No deben estar Tratos Directos, Órdenes de Compra ni Convenio Marco erróneo (CM)
        for excluido in ['CM', 'A2', 'C2', 'D1', 'F2', 'C1', 'F3', 'G1', 'R1', 'SE', 'A1', 'B1', 'E1', 'F1', 'J1']:
            self.assertNotIn(excluido, catalogo['todos_codigos'])

    def test_obtener_catalogo_tipos_licitacion_fallback(self):
        """Verifica que si la base de datos no está disponible, el fallback funcione con los 8 tipos requeridos."""
        from unittest.mock import patch
        from smartbids.utils.licitacion_utils import obtener_catalogo_tipos_licitacion

        with patch('smartbids.models.catalog.TipoLicitacion.objects') as m_tl_objects:
            m_tl_objects.select_related.side_effect = Exception("Error simulado DB")
            catalogo = obtener_catalogo_tipos_licitacion()

        self.assertEqual(len(catalogo['publicas']), 5)
        self.assertEqual(len(catalogo['privadas']), 3)
        self.assertIn('LQ', [p['codigo'] for p in catalogo['publicas']])
        self.assertEqual(next(p for p in catalogo['publicas'] if p['codigo'] == 'LP')['tramo'], '1.000 a 2.000 UTM')
        self.assertEqual(next(p for p in catalogo['publicas'] if p['codigo'] == 'LQ')['tramo'], '2.000 a 5.000 UTM')
        self.assertIn('B2', [p['codigo'] for p in catalogo['privadas']])

    def test_api_catalogos_preferencias_tipos_licitacion(self):
        """Verifica el endpoint /api/catalogos-preferencias/?tipo=tipos_licitacion."""
        import json
        from unittest.mock import patch
        from smartbids.views.perfil_views import catalogos_preferencias

        mock_catalogo = {
            'publicas': [{'codigo': 'L1'}, {'codigo': 'LE'}, {'codigo': 'LP'}, {'codigo': 'LQ'}, {'codigo': 'LR'}],
            'privadas': [{'codigo': 'E2'}, {'codigo': 'CO'}, {'codigo': 'B2'}],
            'todos_codigos': ['L1', 'LE', 'LP', 'LQ', 'LR', 'E2', 'CO', 'B2']
        }

        with patch('smartbids.utils.licitacion_utils.obtener_catalogo_tipos_licitacion', return_value=mock_catalogo):
            req = self.rf.get('/api/catalogos-preferencias/?tipo=tipos_licitacion')
            resp = catalogos_preferencias(req)
            self.assertEqual(resp.status_code, 200)

            data = json.loads(resp.content)
            self.assertEqual(data.get('status'), 'ok')
            self.assertEqual(len(data['datos']['publicas']), 5)
            self.assertEqual(len(data['datos']['privadas']), 3)
            self.assertEqual(data['datos']['publicas'][3]['codigo'], 'LQ')
            self.assertEqual(data['datos']['privadas'][2]['codigo'], 'B2')

    def test_perfil_view_incluye_tipos_procedimiento_en_contexto(self):
        """Verifica que perfil_view pase tipos_procedimiento al contexto de la plantilla."""
        from unittest.mock import patch
        from smartbids.views.page_views import perfil_view

        mock_catalogo = {
            'publicas': [{'codigo': 'L1'}],
            'privadas': [{'codigo': 'B2'}],
            'todos_codigos': ['L1', 'B2']
        }

        with patch('smartbids.views.page_views.obtener_catalogo_tipos_licitacion', return_value=mock_catalogo), \
             patch('smartbids.views.page_views.render') as mock_render:
            mock_render.return_value = 'render_result'
            req = self.rf.get('/perfil/')
            resp = perfil_view(req)
            self.assertEqual(resp, 'render_result')
            mock_render.assert_called_once()
            args, kwargs = mock_render.call_args
            context = args[2] if len(args) > 2 else kwargs.get('context', {})
            self.assertIn('tipos_procedimiento', context)
            self.assertEqual(context['tipos_procedimiento'], mock_catalogo)

    def test_tramos_lp_lq_lr_exactos_en_catalogo_y_plantillas(self):
        """Verifica que LP, LQ y LR reflejen fielmente los tramos de 1.000-2.000, 2.000-5.000 y >5.000 UTM."""
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from smartbids.utils.licitacion_utils import obtener_catalogo_tipos_licitacion

        catalogo = obtener_catalogo_tipos_licitacion()
        publicas_dict = {item['codigo']: item for item in catalogo['publicas']}

        # Verificar LP
        self.assertIn('LP', publicas_dict)
        self.assertEqual(publicas_dict['LP']['tramo'], '1.000 a 2.000 UTM')
        self.assertIn('1.000 UTM e inferior a 2.000 UTM', publicas_dict['LP']['glosa'])

        # Verificar LQ
        self.assertIn('LQ', publicas_dict)
        self.assertEqual(publicas_dict['LQ']['tramo'], '2.000 a 5.000 UTM')
        self.assertIn('2.000 UTM e inferior a 5.000 UTM', publicas_dict['LQ']['glosa'])

        # Verificar LR
        self.assertIn('LR', publicas_dict)
        self.assertEqual(publicas_dict['LR']['tramo'], 'Mayor a 5.000 UTM')

        with patch('smartbids.context_processors.obtener_correos_sistema', return_value={}), \
             patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            # Verificar renderizado de fallback en perfil.html
            req = self.rf.get('/perfil/')
            rendered_perfil = render_to_string('smartbids/perfil.html', {'tipos_procedimiento': {'publicas': [], 'privadas': []}}, request=req)
            self.assertIn('value="LP"', rendered_perfil)
            self.assertIn('1.000 a 2.000 UTM', rendered_perfil)
            self.assertIn('value="LQ"', rendered_perfil)
            self.assertIn('2.000 a 5.000 UTM', rendered_perfil)

            # Verificar renderizado de fallback en Mis_licitaciones.html
            req_mis_lic = self.rf.get('/mis-licitaciones/')
            rendered_mis_lic = render_to_string('smartbids/Mis_licitaciones.html', {'tipos_procedimiento': {'publicas': [], 'privadas': []}, 'estados': []}, request=req_mis_lic)
            self.assertIn('LP – 1.000 a 2.000 UTM', rendered_mis_lic)
            self.assertIn('LQ – 2.000 a 5.000 UTM', rendered_mis_lic)
            self.assertIn('LR – Mayor a 5.000 UTM', rendered_mis_lic)





    def test_mis_licitaciones_view_incluye_tipos_procedimiento_en_contexto(self):
        """Verifica que mis_licitaciones_view pase tipos_procedimiento al contexto de la plantilla."""
        from unittest.mock import patch, MagicMock
        from smartbids.views.licitaciones_views import mis_licitaciones_view

        mock_catalogo = {
            'publicas': [{'codigo': 'L1', 'tramo': 'Menor a 100 UTM'}],
            'privadas': [{'codigo': 'B2', 'tramo': 'Mayor a 1.000 UTM'}],
            'todos_codigos': ['L1', 'B2']
        }

        mock_paginator = MagicMock()
        mock_paginator.get_page.return_value = []

        m_lic_qs = MagicMock()
        m_lic_qs.count.return_value = 50
        m_lic_qs.exclude.return_value.exclude.return_value.values_list.return_value.distinct.return_value.order_by.return_value = ['L1', 'B2']

        m_qs = MagicMock()
        m_qs.aggregate.return_value = {'cierre_proximo': 2, 'convenio_marco': 0}

        with patch('smartbids.views.licitaciones_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.views.licitaciones_views.filtrar_licitaciones', return_value=(m_qs, 0)), \
             patch('smartbids.views.licitaciones_views.Licitacion.objects', m_lic_qs), \
             patch('smartbids.views.licitaciones_views.Paginator', return_value=mock_paginator), \
             patch('smartbids.views.licitaciones_views.obtener_catalogo_tipos_licitacion', return_value=mock_catalogo), \
             patch('smartbids.views.licitaciones_views.render') as mock_render:
            mock_render.return_value = 'render_result'
            req = self.rf.get('/mis-licitaciones/')
            resp = mis_licitaciones_view(req)
            self.assertEqual(resp, 'render_result')
            mock_render.assert_called_once()
            args, kwargs = mock_render.call_args
            context = args[2] if len(args) > 2 else kwargs.get('context', {})
            self.assertIn('tipos_procedimiento', context)
            self.assertEqual(context['tipos_procedimiento'], mock_catalogo)
            self.assertIn('licitaciones_perfil', context)

    def test_mis_licitaciones_view_separa_match_con_perfil_de_procesos_disponibles(self):
        """Verifica que en pestaña Todas, las licitaciones de Match con Perfil se separen de Procesos Disponibles."""
        from unittest.mock import patch, MagicMock
        from django.db.models import Q
        from smartbids.views.licitaciones_views import mis_licitaciones_view

        mock_lic_match = MagicMock(lic_codigo='1075963-374-LE26', lic_tipo_licitacion='LE')
        mock_qs_match = MagicMock()
        mock_qs_match.__iter__.return_value = [mock_lic_match]
        mock_qs_match.__getitem__.return_value = [mock_lic_match]

        mock_qs_general = MagicMock()
        mock_qs_general.exclude.return_value = mock_qs_general
        mock_qs_general.aggregate.return_value = {'cierre_proximo': 0, 'convenio_marco': 0}

        mock_paginator = MagicMock()
        mock_paginator.get_page.return_value = []

        m_lic_objects = MagicMock()
        m_lic_objects.count.return_value = 50
        m_lic_objects.exclude.return_value.exclude.return_value.values_list.return_value.distinct.return_value.order_by.return_value = ['LE']
        m_lic_objects.select_related.return_value.filter.return_value.distinct.return_value.order_by.return_value = mock_qs_match

        req = self.rf.get('/mis-licitaciones/')
        req._q_perfil = Q(lic_tipo_licitacion__in=['LE'])

        with patch('smartbids.views.licitaciones_views.resolver_suscriptor', return_value=MagicMock()), \
             patch('smartbids.views.tracking_views._sincronizar_seguimientos'), \
             patch('smartbids.models.core.SeguimientoLicitacion.objects') as m_seg, \
             patch('smartbids.views.licitaciones_views.Preferencia.objects') as m_pref, \
             patch('smartbids.views.licitaciones_views.filtrar_licitaciones', return_value=(mock_qs_general, 1)), \
             patch('smartbids.views.licitaciones_views.Licitacion.objects', m_lic_objects), \
             patch('smartbids.views.licitaciones_views.Paginator', return_value=mock_paginator), \
             patch('smartbids.views.licitaciones_views.obtener_catalogo_tipos_licitacion', return_value={'publicas': [], 'privadas': []}), \
             patch('smartbids.views.licitaciones_views.estado_licitacion', return_value=('ofertas', 'Recepción de ofertas')), \
             patch('smartbids.views.licitaciones_views.render') as mock_render:
            m_seg.filter.return_value.values_list.return_value = []
            m_pref.filter.return_value.first.return_value = None
            mock_render.return_value = 'render_result'
            resp = mis_licitaciones_view(req)
            self.assertEqual(resp, 'render_result')

            # Debe haber excluido el código de match ('1075963-374-LE26') del queryset general
            mock_qs_general.exclude.assert_called_with(lic_codigo__in=['1075963-374-LE26'])

            context = mock_render.call_args[0][2]
            self.assertEqual(len(context['licitaciones_perfil']), 1)
            self.assertEqual(context['licitaciones_perfil'][0].lic_codigo, '1075963-374-LE26')
            self.assertEqual(context['match_perfil'], 1)

    def test_mis_licitaciones_view_optimizacion_paginador_y_primera_pagina(self):
        """Verifica que el paginador use 10 elementos por página y que en la página 2 no se repitan las tarjetas de perfil."""
        from unittest.mock import patch, MagicMock
        from django.db.models import Q
        from smartbids.views.licitaciones_views import mis_licitaciones_view

        mock_lic_match = MagicMock(lic_codigo='1075963-374-LE26', lic_tipo_licitacion='LE')
        mock_qs_match = MagicMock()
        mock_qs_match.__iter__.return_value = [mock_lic_match]
        mock_qs_match.__getitem__.return_value = [mock_lic_match]

        mock_qs_general = MagicMock()
        mock_qs_general.exclude.return_value = mock_qs_general
        mock_qs_general.aggregate.return_value = {'cierre_proximo': 0, 'convenio_marco': 0}

        m_lic_objects = MagicMock()
        m_lic_objects.count.return_value = 50
        m_lic_objects.exclude.return_value.exclude.return_value.values_list.return_value.distinct.return_value.order_by.return_value = ['LE']
        m_lic_objects.select_related.return_value.filter.return_value.distinct.return_value.order_by.return_value = mock_qs_match

        req_page2 = self.rf.get('/mis-licitaciones/?page=2')
        req_page2._q_perfil = Q(lic_tipo_licitacion__in=['LE'])

        with patch('smartbids.views.licitaciones_views.resolver_suscriptor', return_value=MagicMock()), \
             patch('smartbids.views.tracking_views._sincronizar_seguimientos'), \
             patch('smartbids.models.core.SeguimientoLicitacion.objects') as m_seg, \
             patch('smartbids.views.licitaciones_views.Preferencia.objects') as m_pref, \
             patch('smartbids.views.licitaciones_views.filtrar_licitaciones', return_value=(mock_qs_general, 1)), \
             patch('smartbids.views.licitaciones_views.Licitacion.objects', m_lic_objects), \
             patch('smartbids.views.licitaciones_views.Paginator') as mock_paginator, \
             patch('smartbids.views.licitaciones_views.obtener_catalogo_tipos_licitacion', return_value={'publicas': [], 'privadas': []}), \
             patch('smartbids.views.licitaciones_views.estado_licitacion', return_value=('ofertas', 'Recepción de ofertas')), \
             patch('smartbids.views.licitaciones_views.render') as mock_render:
            mock_paginator.return_value.get_page.return_value = []
            m_seg.filter.return_value.values_list.return_value = []
            m_pref.filter.return_value.first.return_value = None
            mock_render.return_value = 'render_result'

            mis_licitaciones_view(req_page2)

            # Debe instanciar Paginator con 10 items por página
            mock_paginator.assert_called_once_with(mock_qs_general, 10)

            # En la página 2, licitaciones_perfil debe ser lista vacía para no sobrecargar
            context = mock_render.call_args[0][2]
            self.assertEqual(context['licitaciones_perfil'], [])

    def test_cierre_proximo_flag_y_propiedad(self):
        """Verifica que licitaciones que cierran en menos de 3 días marquen es_cierre_proximo en True."""
        from datetime import timedelta
        from django.utils import timezone
        from smartbids.models.procurement import Licitacion

        hoy = timezone.localdate()
        # Licitación cerrando mañana (< 3 días)
        lic_urgente = Licitacion(lic_codigo='URGENTE-01', lic_fecha_cierre=hoy + timedelta(days=1))
        self.assertTrue(lic_urgente.es_cierre_proximo)

        # Licitación cerrando en 10 días (> 3 días)
        lic_lejana = Licitacion(lic_codigo='LEJANA-02', lic_fecha_cierre=hoy + timedelta(days=10))
        self.assertFalse(lic_lejana.es_cierre_proximo)

        # Licitación sin fecha de cierre
        lic_sin_fecha = Licitacion(lic_codigo='SIN-FECHA-03', lic_fecha_cierre=None)
        self.assertFalse(lic_sin_fecha.es_cierre_proximo)

        # Verificación del setter
        lic_lejana.es_cierre_proximo = True
        self.assertTrue(lic_lejana.es_cierre_proximo)
        lic_urgente.es_cierre_proximo = False
        self.assertFalse(lic_urgente.es_cierre_proximo)

    def test_filtrar_licitaciones_tipo_procedimiento_interseccion_and(self):
        """Verifica que el tipo de procedimiento actúe como filtro restrictivo (AND) y no unión (OR)."""
        from unittest.mock import patch, MagicMock
        from django.db.models import Q
        from smartbids.views.licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs

        with patch('smartbids.views.licitaciones_views.Licitacion.objects', mock_qs), \
             patch('smartbids.views.licitaciones_views.Preferencia.objects.filter') as mock_pref:

            # Caso 1: Solo tiene tipo LE seleccionado (como en la captura del usuario)
            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=[],
                pref_tipo_licitacion=['LE'],
                pref_ucom=[],
                pref_palabras_claves=[],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/?tab=claves')
            filtrar_licitaciones(req, suscriptor=MagicMock())

            self.assertTrue(mock_qs.filter.called)
            filter_arg = mock_qs.filter.call_args[0][0]
            self.assertIn(('lic_tipo_licitacion__in', ['LE']), filter_arg.children)

            # Caso 2: Tiene palabras clave Y tipo LE (deben unirse con AND, no con OR)
            mock_qs.reset_mock()
            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=[],
                pref_tipo_licitacion=['LE'],
                pref_ucom=[],
                pref_palabras_claves=['servidor'],
                pref_productos=[]
            )
            req2 = self.rf.get('/mis-licitaciones/?tab=claves')
            filtrar_licitaciones(req2, suscriptor=MagicMock())

            self.assertTrue(mock_qs.filter.called)
            filter_arg2 = mock_qs.filter.call_args[0][0]
            self.assertEqual(filter_arg2.connector, Q.AND)
            self.assertIn(('lic_tipo_licitacion__in', ['LE']), filter_arg2.children)

    def test_filtrar_licitaciones_asigna_q_perfil_en_request(self):
        """Verifica que filtrar_licitaciones guarde _q_perfil en el objeto request."""
        from unittest.mock import patch, MagicMock
        from smartbids.views.licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs

        with patch('smartbids.views.licitaciones_views.Licitacion.objects', mock_qs), \
             patch('smartbids.views.licitaciones_views.Preferencia.objects.filter') as mock_pref:
            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=[],
                pref_tipo_licitacion=['LE'],
                pref_ucom=[],
                pref_palabras_claves=[],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/')
            filtrar_licitaciones(req, suscriptor=MagicMock())
            self.assertTrue(hasattr(req, '_q_perfil'))
            self.assertIn(('lic_tipo_licitacion__in', ['LE']), req._q_perfil.children)

    def test_obtener_catalogo_organismos_ucom_estructura_completa(self):
        """Verifica que obtener_catalogo_organismos_ucom arme la estructura jerárquica Sector -> Organismo -> Unidades."""
        from unittest.mock import patch, MagicMock
        from smartbids.utils.licitacion_utils import obtener_catalogo_organismos_ucom, CACHE_KEY_CATALOGO_UCOM_ARBOL
        from django.core.cache import cache
        cache.delete(CACHE_KEY_CATALOGO_UCOM_ARBOL)

        mock_sectores = [{'codigo_sector': '23', 'nombre_sector': 'Seguridad'}]
        mock_organismos = [
            {'codigo_organismo': 113812, 'org_nombre': 'Carabineros de Chile', 'org_sigla': 'CARAB', 'org_codigo_sector_id': '23'},
            {'codigo_organismo': 111880, 'org_nombre': 'Policía de Investigaciones de Chile', 'org_sigla': 'PDI', 'org_codigo_sector_id': '23'},
        ]
        mock_ucoms = [
            {'codigo_unidad_compra': 3411, 'ucom_descripcion': 'Pref. Bio-Bio', 'ucom_codigo_organismo': 113812},
            {'codigo_unidad_compra': 2939, 'ucom_descripcion': 'Departamento de Apoyo Logistico', 'ucom_codigo_organismo': 111880},
        ]

        with patch('smartbids.utils.licitacion_utils.Sector.objects.values') as m_sec, \
             patch('smartbids.utils.licitacion_utils.Organismo.objects.values') as m_org, \
             patch('smartbids.utils.licitacion_utils.UnidadCompra.objects.values') as m_ucom:
            m_sec.return_value.order_by.return_value = mock_sectores
            m_org.return_value.order_by.return_value = mock_organismos
            m_ucom.return_value.order_by.return_value = mock_ucoms

            res = obtener_catalogo_organismos_ucom()
            self.assertEqual(res['status'], 'ok')
            self.assertEqual(len(res['sectores']), 1)
            sec = res['sectores'][0]
            self.assertEqual(sec['codigo_sector'], '23')
            self.assertEqual(sec['nombre_sector'], 'Seguridad')
            self.assertEqual(len(sec['organismos']), 2)

            # Verificar Carabineros y su unidad
            carab = next(o for o in sec['organismos'] if o['codigo_organismo'] == '113812')
            self.assertEqual(carab['nombre_organismo'], 'Carabineros de Chile')
            self.assertEqual(len(carab['unidades']), 1)
            self.assertEqual(carab['unidades'][0]['codigo_ucom'], 3411)

            # Verificar PDI y su unidad
            pdi = next(o for o in sec['organismos'] if o['codigo_organismo'] == '111880')
            self.assertEqual(pdi['nombre_organismo'], 'Policía de Investigaciones de Chile')
            self.assertEqual(len(pdi['unidades']), 1)
            self.assertEqual(pdi['unidades'][0]['codigo_ucom'], 2939)

            cache.delete(CACHE_KEY_CATALOGO_UCOM_ARBOL)

    def test_obtener_catalogo_organismos_ucom_filtro_sector_y_q(self):
        """Verifica el filtrado por sector y búsqueda textual por nombre o sigla."""
        from unittest.mock import patch, MagicMock
        from smartbids.utils.licitacion_utils import obtener_catalogo_organismos_ucom, CACHE_KEY_CATALOGO_UCOM_ARBOL
        from django.core.cache import cache
        cache.delete(CACHE_KEY_CATALOGO_UCOM_ARBOL)

        mock_sectores = [
            {'codigo_sector': '18', 'nombre_sector': 'Municipio'},
            {'codigo_sector': '23', 'nombre_sector': 'Seguridad'}
        ]
        mock_organismos = [
            {'codigo_organismo': 87399, 'org_nombre': 'Ilustre Municipalidad de Viña del Mar', 'org_sigla': 'VIÑA', 'org_codigo_sector_id': '18'},
            {'codigo_organismo': 113812, 'org_nombre': 'Carabineros de Chile', 'org_sigla': 'CARAB', 'org_codigo_sector_id': '23'},
        ]
        mock_ucoms = [
            {'codigo_unidad_compra': 1234, 'ucom_descripcion': 'Secpla Viña', 'ucom_codigo_organismo': 87399},
            {'codigo_unidad_compra': 3411, 'ucom_descripcion': 'Pref. Bio-Bio', 'ucom_codigo_organismo': 113812},
        ]

        with patch('smartbids.utils.licitacion_utils.Sector.objects.values') as m_sec, \
             patch('smartbids.utils.licitacion_utils.Organismo.objects.values') as m_org, \
             patch('smartbids.utils.licitacion_utils.UnidadCompra.objects.values') as m_ucom:
            m_sec.return_value.order_by.return_value = mock_sectores
            m_org.return_value.order_by.return_value = mock_organismos
            m_ucom.return_value.order_by.return_value = mock_ucoms

            # Filtrar por sector 18 (Municipio)
            res_sec = obtener_catalogo_organismos_ucom(sector_codigo='18')
            self.assertEqual(len(res_sec['sectores']), 1)
            self.assertEqual(res_sec['sectores'][0]['codigo_sector'], '18')

            # Buscar 'viña'
            res_q = obtener_catalogo_organismos_ucom(q='viña')
            self.assertEqual(len(res_q['sectores']), 1)
            self.assertEqual(res_q['sectores'][0]['organismos'][0]['nombre_organismo'], 'Ilustre Municipalidad de Viña del Mar')

            # Buscar 'carabineros'
            res_carab = obtener_catalogo_organismos_ucom(q='carabineros')
            self.assertEqual(len(res_carab['sectores']), 1)
            self.assertEqual(res_carab['sectores'][0]['organismos'][0]['nombre_organismo'], 'Carabineros de Chile')

            cache.delete(CACHE_KEY_CATALOGO_UCOM_ARBOL)

    def test_catalogos_preferencias_view_ucom_arbol(self):
        """Verifica la respuesta JSON del endpoint /api/catalogos-preferencias/?tipo=ucom_arbol."""
        from unittest.mock import patch
        import json
        from smartbids.views.perfil_views import catalogos_preferencias

        mock_catalogo = {
            'status': 'ok',
            'sectores': [{'codigo_sector': '23', 'nombre_sector': 'Seguridad', 'organismos': []}],
            'sectores_lista': [{'codigo_sector': '23', 'nombre_sector': 'Seguridad', 'organismos_count': 1}],
            'total_sectores': 1,
            'total_organismos': 1,
            'total_unidades': 1,
        }

        with patch('smartbids.utils.licitacion_utils.obtener_catalogo_organismos_ucom', return_value=mock_catalogo):
            req = self.rf.get('/api/catalogos-preferencias/?tipo=ucom_arbol&sector=23')
            resp = catalogos_preferencias(req)
            self.assertEqual(resp.status_code, 200)
            data = json.loads(resp.content)
            self.assertEqual(data['status'], 'ok')
            self.assertIn('sectores', data)
            self.assertIn('sectores_lista', data)

    def test_filtrar_licitaciones_ucom_ids_incluye_organismo(self):
        """Verifica que ucom_ids en el perfil aplique OR entre lic_codigo_ucom_id y ucom_codigo_organismo."""
        from unittest.mock import patch, MagicMock
        from django.db.models import Q
        from smartbids.views.licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs

        with patch('smartbids.views.licitaciones_views.Licitacion.objects', mock_qs), \
             patch('smartbids.views.licitaciones_views.Preferencia.objects.filter') as mock_pref:
            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=[],
                pref_tipo_licitacion=[],
                pref_ucom=[113812],
                pref_palabras_claves=[],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/?tab=claves')
            filtrar_licitaciones(req, suscriptor=MagicMock())

            self.assertTrue(mock_qs.filter.called)
            filter_arg = mock_qs.filter.call_args[0][0]
            # Debe contener la condición para lic_codigo_ucom_id o ucom_codigo_organismo
            children_keys = [c[0] for c in filter_arg.children if isinstance(c, tuple)]
            self.assertTrue('lic_codigo_ucom_id__in' in children_keys or any('ucom_codigo_organismo' in str(c) for c in filter_arg.children))


class DetalleLicitacionPrioridadTests(SimpleTestCase):
    def setUp(self):
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def tearDown(self):
        from django.core.cache import cache
        cache.clear()

    def test_detalle_licitacion_prioriza_datos_db_completos_sin_llamar_mercado_publico(self):
        """Verifica que si la BD ya tiene descripción, documentos y preguntas, no llame a Mercado Público."""
        from unittest.mock import patch
        from types import SimpleNamespace
        from smartbids.views.tracking_views import detalle_licitacion_view

        req = self.rf.get('/licitaciones/1000-1-LE26/')

        ucom = SimpleNamespace(
            ucom_descripcion='Hospital Base',
            ucom_direccion='Calle Principal 123',
            ucom_rut='70.000.000-1',
            ucom_codigo_comuna='13101',
        )
        lic = SimpleNamespace(
            lic_codigo='1000-1-LE26',
            lic_nombre='Servicio de Mantención',
            lic_descripcion='Descripción completa almacenada en BD',
            lic_texto='Texto alternativo',
            lic_fecha_public=None,
            lic_fecha_preguntas=None,
            lic_fecha_respuestas=None,
            lic_fecha_cierre=None,
            lic_codigo_ucom=ucom,
            lic_url='https://www.mercadopublico.cl/fichaLicitacion.html?idLicitacion=1000-1-LE26',
            lic_tipo_licitacion='LE',
            lic_presupuesto=1000000,
            lic_moneda='CLP',
            lic_items_cantidad=3,
        )

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None), \
             patch('smartbids.views.tracking_views.get_object_or_404', return_value=lic), \
             patch('smartbids.views.tracking_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.views.tracking_views.LicitacionDocumento.objects.filter') as mock_docs, \
             patch('smartbids.views.tracking_views.LicitacionPregunta.objects.filter') as mock_pregs, \
             patch('smartbids.views.tracking_views.obtener_datos_externos_licitacion') as mock_ext:

            mock_docs.return_value = [
                SimpleNamespace(nombre='Bases.pdf', url='http://bases.pdf', extension='pdf', descripcion='Bases', fecha=None)
            ]
            mock_pregs.return_value = [
                SimpleNamespace(pregunta='¿Plazo?', respuesta='30 días', fecha_pregunta=None, fecha_respuesta=None)
            ]

            resp = detalle_licitacion_view(req, '1000-1-LE26')
            self.assertEqual(resp.status_code, 200)
            # Como la BD tenía toda la información, no debe llamar a Mercado Público
            self.assertFalse(mock_ext.called)
            self.assertContains(resp, 'Descripción completa almacenada en BD')

    def test_detalle_licitacion_llama_mercado_publico_si_falta_descripcion(self):
        """Verifica que si la BD NO tiene descripción, se llame a Mercado Público y se actualice en BD."""
        from unittest.mock import patch
        from types import SimpleNamespace
        from smartbids.views.tracking_views import detalle_licitacion_view

        req = self.rf.get('/licitaciones/1000-2-LE26/')

        lic = SimpleNamespace(
            lic_codigo='1000-2-LE26',
            lic_nombre='Adquisición Insumos',
            lic_descripcion='',  # Vacía en BD
            lic_texto='',        # Vacía en BD
            lic_fecha_public=None,
            lic_fecha_preguntas=None,
            lic_fecha_respuestas=None,
            lic_fecha_cierre=None,
            lic_codigo_ucom=None,
            lic_url='',
            lic_tipo_licitacion='LE',
            lic_presupuesto=500000,
            lic_moneda='CLP',
            lic_items_cantidad=1,
        )

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None), \
             patch('smartbids.views.tracking_views.get_object_or_404', return_value=lic), \
             patch('smartbids.views.tracking_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.views.tracking_views.LicitacionDocumento.objects.filter', return_value=[]), \
             patch('smartbids.views.tracking_views.LicitacionPregunta.objects.filter', return_value=[]), \
             patch('smartbids.views.tracking_views.Licitacion.objects.filter') as mock_lic_filter, \
             patch('smartbids.views.tracking_views.obtener_datos_externos_licitacion') as mock_ext:

            mock_ext.return_value = {
                'descripcion': 'Descripción recuperada desde la ficha oficial de Mercado Público',
                'url_adjuntos': 'https://www.mercadopublico.cl/adjuntos',
                'url_preguntas': 'https://www.mercadopublico.cl/preguntas',
                'preguntas': [],
                'direccion': 'Dirección externa',
            }

            resp = detalle_licitacion_view(req, '1000-2-LE26')
            self.assertEqual(resp.status_code, 200)
            # Verificamos que se haya llamado a Mercado Público porque faltaba la descripción
            self.assertTrue(mock_ext.called)
            # Verificamos que se haya intentado actualizar la BD con la descripción recuperada
            mock_lic_filter.return_value.update.assert_called_with(
                lic_descripcion='Descripción recuperada desde la ficha oficial de Mercado Público'
            )
            self.assertContains(resp, 'Descripción recuperada desde la ficha oficial de Mercado Público')

    def test_detalle_licitacion_preserva_descripcion_db_al_consultar_adjuntos_externos(self):
        """Verifica que si la BD ya tiene descripción pero faltan adjuntos, mantenga la descripción de la BD."""
        from unittest.mock import patch
        from types import SimpleNamespace
        from smartbids.views.tracking_views import detalle_licitacion_view

        req = self.rf.get('/licitaciones/1000-3-LE26/')

        lic = SimpleNamespace(
            lic_codigo='1000-3-LE26',
            lic_nombre='Servicio Especializado',
            lic_descripcion='Descripción propia y única de la Base de Datos',
            lic_texto='Texto alternativo',
            lic_fecha_public=None,
            lic_fecha_preguntas=None,
            lic_fecha_respuestas=None,
            lic_fecha_cierre=None,
            lic_codigo_ucom=None,
            lic_url='',
            lic_tipo_licitacion='LE',
            lic_presupuesto=100000,
            lic_moneda='CLP',
            lic_items_cantidad=1,
        )

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None), \
             patch('smartbids.views.tracking_views.get_object_or_404', return_value=lic), \
             patch('smartbids.views.tracking_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.views.tracking_views.LicitacionDocumento.objects.filter', return_value=[]), \
             patch('smartbids.views.tracking_views.LicitacionPregunta.objects.filter', return_value=[]), \
             patch('smartbids.views.tracking_views.Licitacion.objects.filter') as mock_lic_filter, \
             patch('smartbids.views.tracking_views.obtener_datos_externos_licitacion') as mock_ext:

            mock_ext.return_value = {
                'descripcion': 'Otra descripcion de mercado publico que NO debe sobreescribir la BD',
                'url_adjuntos': 'https://www.mercadopublico.cl/adjuntos-oficiales',
                'url_preguntas': 'https://www.mercadopublico.cl/preguntas-oficiales',
                'preguntas': [],
                'direccion': 'Dirección externa',
            }

            resp = detalle_licitacion_view(req, '1000-3-LE26')
            self.assertEqual(resp.status_code, 200)
            self.assertTrue(mock_ext.called)
            # Como la BD ya tenía descripción, NUNCA debe sobreescribirla en la base de datos
            self.assertFalse(mock_lic_filter.return_value.update.called)
            # La respuesta debe contener la descripción de la BD
            self.assertContains(resp, 'Descripción propia y única de la Base de Datos')
            self.assertNotContains(resp, 'Otra descripcion de mercado publico que NO debe sobreescribir la BD')
            # Y debe contener el enlace de adjuntos externos rescatado
            self.assertContains(resp, 'https://www.mercadopublico.cl/adjuntos-oficiales')


class Seguridad2FAYCredencialesTests(SimpleTestCase):
    def setUp(self):
        import json
        from django.test import RequestFactory
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def test_enviar_codigo_login_no_expone_otp_en_respuesta_json(self):
        """Verifica que enviar_codigo_login no exponga el código OTP ni contraseñas en el JSON devuelto."""
        import json
        from unittest.mock import patch
        from smartbids.views.auth_views import enviar_codigo_login, CODIGOS_OTP_TEMPORALES

        req = self.rf.post(
            '/api/enviar-codigo-login/',
            data=json.dumps({'email': 'usuario.seguro@smartbids.cl'}).encode('utf-8'),
            content_type='application/json'
        )

        with patch('smartbids.views.auth_views.send_mail') as mock_send_mail:
            resp = enviar_codigo_login(req)
            self.assertEqual(resp.status_code, 200)
            data = json.loads(resp.content.decode('utf-8'))

            # La respuesta solo debe indicar el estado ok y mensaje de éxito
            self.assertEqual(data.get('status'), 'ok')
            self.assertIn('mensaje', data)

            # NUNCA debe contener la clave 'codigo', 'otp' ni 'password'
            self.assertNotIn('codigo', data)
            self.assertNotIn('otp', data)
            self.assertNotIn('password', data)

            # El código debe almacenarse de forma segura (hasheado con salt)
            self.assertIn('usuario.seguro@smartbids.cl', CODIGOS_OTP_TEMPORALES)
            registro = CODIGOS_OTP_TEMPORALES['usuario.seguro@smartbids.cl']
            self.assertIn('hash', registro)
            self.assertNotIn('codigo', registro)  # Nunca en texto plano

            # El correo debe haberse despachado con el código en el cuerpo del email
            self.assertTrue(mock_send_mail.called)
            cuerpo_mail = mock_send_mail.call_args[1]['message']
            self.assertIn('Tu código de verificación para SmartBids es:', cuerpo_mail)

    def test_validar_codigo_login_flujo_seguridad_y_eliminacion_otp(self):
        """Verifica que validar_codigo_login elimine el OTP inmediatamente para prevenir ataques de repetición."""
        import json
        from unittest.mock import patch, MagicMock
        from django.utils import timezone
        from smartbids.views.auth_views import validar_codigo_login, CODIGOS_OTP_TEMPORALES, _calcular_hash_otp

        email_prueba = 'usuario.otp@smartbids.cl'
        CODIGOS_OTP_TEMPORALES[email_prueba] = {
            'hash': _calcular_hash_otp(email_prueba, '987654'),
            'creado_el': timezone.now(),
            'intentos_fallidos': 0
        }

        mock_susc = MagicMock()

        req = self.rf.post(
            '/api/validar-codigo-login/',
            data=json.dumps({
                'email': email_prueba,
                'uid': 'FIREBASE_UID_TEST_123',
                'codigo': '987654',
                'session_token': 'TOKEN_SESION_ABC'
            }).encode('utf-8'),
            content_type='application/json'
        )

        with patch('smartbids.views.auth_views.Suscriptor.objects.get_or_create', return_value=(mock_susc, False)), \
             patch('smartbids.views.auth_views.EstadoSuscriptor.objects.filter'):
            resp = validar_codigo_login(req)
            self.assertEqual(resp.status_code, 200)
            data = json.loads(resp.content.decode('utf-8'))

            self.assertEqual(data.get('status'), 'ok')
            self.assertNotIn('codigo', data)
            self.assertNotIn('password', data)

            # El token de sesión debe haberse asignado al suscriptor
            self.assertEqual(mock_susc.token_sesion, 'TOKEN_SESION_ABC')
            self.assertTrue(mock_susc.save.called)

            # El código OTP debe haberse eliminado de la memoria para evitar reuso
            self.assertNotIn(email_prueba, CODIGOS_OTP_TEMPORALES)

            # Un segundo intento con el mismo código debe ser rechazado
            req_reintento = self.rf.post(
                '/api/validar-codigo-login/',
                data=json.dumps({
                    'email': email_prueba,
                    'uid': 'FIREBASE_UID_TEST_123',
                    'codigo': '987654',
                    'session_token': 'TOKEN_SESION_ABC'
                }).encode('utf-8'),
                content_type='application/json'
            )
            resp_reintento = validar_codigo_login(req_reintento)
            self.assertEqual(resp_reintento.status_code, 400)

    def test_purgar_codigos_otp_expirados(self):
        """Verifica que purgar_codigos_otp_expirados elimine los códigos que superaron los 10 minutos."""
        from datetime import timedelta
        from django.utils import timezone
        from smartbids.views.auth_views import purgar_codigos_otp_expirados, CODIGOS_OTP_TEMPORALES, revocar_codigo_otp

        ahora = timezone.now()
        CODIGOS_OTP_TEMPORALES['expirado@smartbids.cl'] = {
            'hash': 'hash1',
            'creado_el': ahora - timedelta(seconds=601),
            'intentos_fallidos': 0
        }
        CODIGOS_OTP_TEMPORALES['vigente@smartbids.cl'] = {
            'hash': 'hash2',
            'creado_el': ahora - timedelta(seconds=300),
            'intentos_fallidos': 0
        }

        eliminados = purgar_codigos_otp_expirados()
        self.assertEqual(eliminados, 1)
        self.assertNotIn('expirado@smartbids.cl', CODIGOS_OTP_TEMPORALES)
        self.assertIn('vigente@smartbids.cl', CODIGOS_OTP_TEMPORALES)

        # Probar revocación explícita
        self.assertTrue(revocar_codigo_otp('vigente@smartbids.cl'))
        self.assertNotIn('vigente@smartbids.cl', CODIGOS_OTP_TEMPORALES)

    def test_validar_codigo_login_limite_intentos_fallidos_bloqueo(self):
        """Verifica que tras 5 intentos fallidos el OTP sea destruido por protección contra fuerza bruta."""
        import json
        from django.utils import timezone
        from smartbids.views.auth_views import validar_codigo_login, CODIGOS_OTP_TEMPORALES, _calcular_hash_otp

        email = 'bruteforce@smartbids.cl'
        CODIGOS_OTP_TEMPORALES[email] = {
            'hash': _calcular_hash_otp(email, '123456'),
            'creado_el': timezone.now(),
            'intentos_fallidos': 4
        }

        # Intento número 5 (debe fallar y destruir el código)
        req = self.rf.post(
            '/api/validar-codigo-login/',
            data=json.dumps({
                'email': email,
                'uid': 'UID_TEST',
                'codigo': '000000',
                'session_token': 'TOKEN_TEST'
            }).encode('utf-8'),
            content_type='application/json'
        )
        resp = validar_codigo_login(req)
        self.assertEqual(resp.status_code, 429)
        self.assertNotIn(email, CODIGOS_OTP_TEMPORALES)

    def test_archivos_javascript_no_almacenan_passwords_en_variables_globales(self):
        """Verifica estáticamente que los archivos JavaScript no almacenen variables pendingPassword."""
        import os

        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        js_dir = os.path.join(base_dir, 'static', 'js')

        archivos_analizados = 0
        for root, _, files in os.walk(js_dir):
            for file in files:
                if file.endswith('.js'):
                    path_file = os.path.join(root, file)
                    with open(path_file, 'r', encoding='utf-8') as f:
                        contenido = f.read()
                        self.assertNotIn(
                            'pendingPassword',
                            contenido,
                            f'Se detectó variable de contraseña en JavaScript en {file}'
                        )
                    archivos_analizados += 1

        self.assertGreater(archivos_analizados, 0, 'Deben haberse analizado archivos JS.')

    def test_obtener_nombre_comuna_resuelve_correctamente(self):
        """Verifica que obtener_nombre_comuna traduzca códigos a nombres legibles."""
        from smartbids.utils.licitacion_utils import obtener_nombre_comuna
        from unittest.mock import patch

        mapa_mock = {'13101': 'Santiago', '01101': 'Iquique'}
        with patch('smartbids.utils.licitacion_utils.obtener_mapa_comunas', return_value=mapa_mock):
            self.assertEqual(obtener_nombre_comuna('13101'), 'Santiago')
            self.assertEqual(obtener_nombre_comuna('1101'), 'Iquique')
            self.assertEqual(obtener_nombre_comuna('99999'), '99999')
            self.assertEqual(obtener_nombre_comuna(''), '')
            self.assertEqual(obtener_nombre_comuna(None), '')

    def test_propiedad_nombre_comuna_en_unidad_compra_y_licitacion(self):
        """Verifica que las propiedades nombre_comuna en UnidadCompra y Licitacion funcionen."""
        from smartbids.models.procurement import UnidadCompra, Licitacion
        from unittest.mock import patch

        mapa_mock = {'13101': 'Santiago'}
        with patch('smartbids.utils.licitacion_utils.obtener_mapa_comunas', return_value=mapa_mock):
            ucom = UnidadCompra(codigo_unidad_compra=1, ucom_codigo_comuna='13101')
            self.assertEqual(ucom.nombre_comuna, 'Santiago')
            self.assertEqual(ucom.ucom_nombre_comuna, 'Santiago')

            lic = Licitacion(lic_codigo='123-1-L24', lic_codigo_ucom=ucom)
            self.assertEqual(lic.nombre_comuna, 'Santiago')

            lic_sin_ucom = Licitacion(lic_codigo='124-1-L24', lic_codigo_ucom=None)
            self.assertEqual(lic_sin_ucom.nombre_comuna, '')

    def test_filtro_template_nombre_comuna(self):
        """Verifica que el filtro de template nombre_comuna traduzca el código."""
        from smartbids.templatetags.montos import nombre_comuna
        from unittest.mock import patch

        mapa_mock = {'13101': 'Santiago'}
        with patch('smartbids.utils.licitacion_utils.obtener_mapa_comunas', return_value=mapa_mock):
            self.assertEqual(nombre_comuna('13101'), 'Santiago')
            self.assertEqual(nombre_comuna(''), '')
            self.assertEqual(nombre_comuna(None), '')


class LegalComplianceArcopTests(SimpleTestCase):
    """Pruebas de cumplimiento de la Ley N° 21.719, Derechos ARCOP y Comercio Electrónico."""

    def test_footer_renderiza_razon_social_y_rut_placeholder(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/')
        with patch('smartbids.context_processors.obtener_correos_sistema', return_value={}), \
             patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('smartbids/includes/footer.html', {'app_version': 'v1.0.0'}, request=req)

        self.assertIn('no informada', html)
        self.assertIn('00000000-5', html)
        self.assertIn('Política de Cookies', html)
        self.assertIn('Ley 21.719', html)

    def test_registro_template_incluye_casillas_consentimiento(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/registro/')
        with patch('smartbids.context_processors.obtener_correos_sistema', return_value={}), \
             patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('registro/registro.html', request=req)

        self.assertIn('consent-terms', html)
        self.assertIn('consent-marketing', html)
        self.assertIn('Términos del Servicio', html)
        self.assertIn('Política de Privacidad', html)
        self.assertIn('Ley N° 21.719', html)

    def test_informacion_template_cita_ley_21719_y_apdp(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/informacion/')
        with patch('smartbids.context_processors.obtener_correos_sistema', return_value={}), \
             patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('smartbids/informacion.html', request=req)

        self.assertIn('Ley N° 21.719', html)
        self.assertIn('APDP', html)
        self.assertIn('Portabilidad', html)
        self.assertIn('no informada', html)
        self.assertIn('00000000-5', html)

    def test_cookie_banner_presente_en_base(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/')
        with patch('smartbids.context_processors.obtener_correos_sistema', return_value={}), \
             patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('smartbids/base.html', request=req)

        self.assertIn('smartbids-cookie-banner', html)
        self.assertIn('btn-accept-cookies', html)
        self.assertIn('cookie-banner.js', html)

    def test_exportar_datos_arcop_exito(self):
        import json
        from unittest.mock import patch, MagicMock
        from django.test import RequestFactory
        from smartbids.views.perfil_views import exportar_datos_arcop

        rf = RequestFactory()
        req = rf.post(
            '/api/arcop/exportar-datos/',
            data=json.dumps({'uid': 'test-uid-123'}),
            content_type='application/json'
        )

        mock_suscriptor = MagicMock()
        mock_suscriptor.id_suscriptor = 99
        mock_suscriptor.firebase_uid = 'test-uid-123'
        mock_suscriptor.sus_nombre1 = 'Juan'
        mock_suscriptor.sus_nombre2 = ''
        mock_suscriptor.sus_apellido1 = 'Pérez'
        mock_suscriptor.sus_apellido2 = ''
        mock_suscriptor.sus_nombre_social = ''
        mock_suscriptor.sus_iniciales = 'JP'
        mock_suscriptor.codigo_estado.nombre_estado = 'Habilitado'
        mock_suscriptor.fecha_registro = None
        mock_suscriptor.fecha_actualizacion = None
        mock_suscriptor.sus_rut_empresa = None
        mock_suscriptor.consentimiento_terminos_fecha = None
        mock_suscriptor.consentimiento_terminos_version = 'v1.0'
        mock_suscriptor.consentimiento_marketing = False

        with patch('smartbids.views.perfil_views.Suscriptor.objects.select_related') as mock_sr, \
             patch('smartbids.views.perfil_views.validar_autorizacion_suscriptor', return_value=(True, '')), \
             patch('smartbids.views.perfil_views.Preferencia.objects.filter') as mock_pref, \
             patch('smartbids.models.core.SolicitudArcop.objects.create') as mock_arcop_log:

            mock_sr.return_value.filter.return_value.first.return_value = mock_suscriptor
            mock_pref.return_value.first.return_value = None

            resp = exportar_datos_arcop(req)

        self.assertEqual(resp.status_code, 200)
        data = json.loads(resp.content.decode('utf-8'))
        self.assertIn('declaracion_legal', data)
        self.assertIn('titular', data)
        self.assertEqual(data['titular']['firebase_uid'], 'test-uid-123')
        self.assertIn('Ley N° 21.719', data['declaracion_legal']['normativa'])

    def test_canal_privacidad_se_lee_dinamicamente_de_info_contacto(self):
        """Verifica que el canal de privacidad, Razón Social y RUT se lean dinámicamente de Contacto y Soporte."""
        from unittest.mock import patch, MagicMock
        from django.template.loader import render_to_string
        from django.test import RequestFactory
        from smartbids.context_processors import footer_info_context

        rf = RequestFactory()
        req = rf.get('/informacion/')

        mock_info = MagicMock()
        mock_info.contacto_texto = (
            "Agua Santa 1400\n"
            "2520000 Viña del Mar\n"
            "Mobile: +56 9 6595 6863\n"
            "Email: smartbids.qa@chileavanza.cl\n"
            "Razón Social: Mi Empresa SpA\n"
            "RUT: 76.543.210-K"
        )
        mock_info.empresas_enlaces = []

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=mock_info):
            ctx = footer_info_context(req)
            self.assertEqual(ctx['contacto_email'], 'smartbids.qa@chileavanza.cl')
            self.assertEqual(ctx['razon_social'], 'Mi Empresa SpA')
            self.assertEqual(ctx['rut'], '76.543.210-K')

            html_info = render_to_string('smartbids/informacion.html', ctx, request=req)
            self.assertIn('smartbids.qa@chileavanza.cl', html_info)
            self.assertIn('Mi Empresa SpA', html_info)
            self.assertIn('76.543.210-K', html_info)

            html_perfil = render_to_string('smartbids/perfil.html', ctx, request=req)
            self.assertIn('smartbids.qa@chileavanza.cl', html_perfil)

            html_footer = render_to_string('smartbids/includes/footer.html', ctx, request=req)
            self.assertIn('Mi Empresa SpA', html_footer)
            self.assertIn('76.543.210-K', html_footer)
























