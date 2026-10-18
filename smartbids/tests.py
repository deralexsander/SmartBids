import csv
from datetime import date, datetime, time, timezone as datetime_timezone
from io import StringIO

from django.test import SimpleTestCase, override_settings

from .licitacion_utils import estado_licitacion, escribir_fila_csv
from .models.procurement import Licitacion
from .templatetags.montos import miles_chilenos


@override_settings(TIME_ZONE='UTC')
class EstadoLicitacionTests(SimpleTestCase):
    ahora = datetime(2026, 10, 13, 12, 0, tzinfo=datetime_timezone.utc)

    def test_calcula_periodo_de_preguntas(self):
        licitacion = Licitacion(
            lic_fecha_cierre=date(2026, 10, 20),
            lic_fecha_preguntas=datetime(2026, 10, 13, 11, 0, tzinfo=datetime_timezone.utc),
            lic_fecha_respuestas=datetime(2026, 10, 14, 11, 0, tzinfo=datetime_timezone.utc),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora), ('preguntas', 'Período de preguntas'))

    def test_calcula_recepcion_de_ofertas(self):
        licitacion = Licitacion(
            lic_fecha_cierre=date(2026, 10, 20),
            lic_fecha_respuestas=datetime(2026, 10, 13, 11, 0, tzinfo=datetime_timezone.utc),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora)[0], 'ofertas')

    def test_acepta_fechas_legacy_sin_zona_horaria(self):
        licitacion = Licitacion(
            lic_fecha_cierre=date(2026, 10, 20),
            lic_fecha_preguntas=datetime(2026, 10, 13, 11, 0),
            lic_fecha_respuestas=datetime(2026, 10, 14, 11, 0),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora)[0], 'preguntas')

    def test_cierre_tiene_precedencia_sobre_otros_hitos(self):
        licitacion = Licitacion(
            lic_fecha_cierre=date(2026, 10, 12),
            lic_fecha_respuestas=datetime(2026, 10, 13, 11, 0, tzinfo=datetime_timezone.utc),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora)[0], 'cerrada')

    def test_estado_inicial_si_no_hay_hitos_vencidos(self):
        licitacion = Licitacion(
            lic_fecha_public=date(2026, 10, 13),
            lic_fecha_cierre=date(2026, 10, 20),
            lic_hora_cierre=time(17, 0),
        )
        self.assertEqual(estado_licitacion(licitacion, self.ahora)[0], 'publicada')

    def test_exportacion_csv_protege_formulas_de_hoja_de_calculo(self):
        output = StringIO()
        writer = csv.writer(output)
        escribir_fila_csv(writer, ('=1+1', 'Texto normal'))
        self.assertEqual(next(csv.reader(StringIO(output.getvalue()))), ["'=1+1", 'Texto normal'])

    def test_exportacion_excel_genera_archivo_y_protege_formulas(self):
        from io import BytesIO
        import openpyxl
        from .licitacion_utils import generar_excel_response

        response = generar_excel_response(
            'test.xlsx',
            ('Columna 1', 'Columna 2'),
            [('=SUM(A1:A2)', 'Texto simple'), ('@link', 123)]
        )
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('test.xlsx', response['Content-Disposition'])

        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, 'Columna 1')
        self.assertEqual(ws.cell(row=2, column=1).value, "'=SUM(A1:A2)")
        self.assertEqual(ws.cell(row=2, column=2).value, 'Texto simple')
        self.assertEqual(ws.cell(row=3, column=1).value, "'@link")
        self.assertEqual(ws.cell(row=3, column=2).value, 123)


class FormatoMontosTests(SimpleTestCase):
    def test_formatea_montos_con_separador_de_miles_chileno(self):
        self.assertEqual(miles_chilenos('125000000'), '125.000.000')

    def test_redondea_montos_sin_decimales(self):
        self.assertEqual(miles_chilenos('125000000.50'), '125.000.001')


class PreferenciasTests(SimpleTestCase):
    def test_normalizar_cadena_simple(self):
        from .perfil_views import normalizar_cadena_simple
        self.assertEqual(normalizar_cadena_simple("Mantención"), "mantencion")
        self.assertEqual(normalizar_cadena_simple("¡Cloud & Seguridad!"), "cloud seguridad")
        self.assertEqual(normalizar_cadena_simple(""), "")


class VistasExportacionTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        self.rf = RequestFactory()

    def test_exportar_licitaciones_csv_por_defecto(self):
        from unittest.mock import patch
        from .licitaciones_views import exportar_licitaciones_view

        with patch('smartbids.licitaciones_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.licitaciones_views.filtrar_licitaciones', return_value=([], 0)):
            req = self.rf.get('/mis-licitaciones/exportar/')
            resp = exportar_licitaciones_view(req)
            self.assertEqual(resp['Content-Type'], 'text/csv; charset=utf-8')
            self.assertIn('licitaciones-smartbids.csv', resp['Content-Disposition'])

    def test_exportar_licitaciones_excel(self):
        from unittest.mock import patch
        from .licitaciones_views import exportar_licitaciones_view

        with patch('smartbids.licitaciones_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.licitaciones_views.filtrar_licitaciones', return_value=([], 0)):
            req = self.rf.get('/mis-licitaciones/exportar/?formato=xlsx')
            resp = exportar_licitaciones_view(req)
            self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            self.assertIn('licitaciones-smartbids.xlsx', resp['Content-Disposition'])

    def test_exportar_dashboard_excel(self):
        from unittest.mock import patch
        from .page_views import exportar_dashboard_view

        req = self.rf.get('/dashboard/exportar/?formato=xlsx&gran=mes&periodo=2024-05')
        resp = exportar_dashboard_view(req)
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('dashboard-tablas-2024-05.xlsx', resp['Content-Disposition'])

    def test_exportar_dashboard_csv(self):
        from .page_views import exportar_dashboard_view

        req = self.rf.get('/dashboard/exportar/?formato=csv&gran=mes&periodo=2024-05')
        resp = exportar_dashboard_view(req)
        self.assertEqual(resp['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('dashboard-tablas-2024-05.csv', resp['Content-Disposition'])
        self.assertIn('=== RESUMEN KPIS ===', resp.content.decode('utf-8'))



class DashboardDWTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def test_dashboard_page_render(self):
        from unittest.mock import patch
        from .dashboard_views import dashboard_view
        req = self.rf.get('/dashboard/')
        with patch('smartbids.dashboard_views.render') as mock_render:
            mock_render.return_value = 'ok'
            resp = dashboard_view(req)
            self.assertEqual(resp, 'ok')
            mock_render.assert_called_once_with(req, 'smartbids/dashboard.html')

    def test_dashboard_periodo_invalido(self):
        from .dashboard_views import dashboard_kpis
        req = self.rf.get('/dashboard/kpis/?gran=mensual&periodo=invalido')
        resp = dashboard_kpis(req)
        self.assertEqual(resp.status_code, 400)


class ReenvioActivacionTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        from django.core.cache import cache
        self.rf = RequestFactory()
        cache.clear()

    def test_metodo_no_permitido(self):
        from .auth_views import reenviar_correo_activacion
        req = self.rf.get('/api/reenviar-correo-activacion/')
        resp = reenviar_correo_activacion(req)
        self.assertEqual(resp.status_code, 405)

    def test_correo_requerido(self):
        from .auth_views import reenviar_correo_activacion
        req = self.rf.post('/api/reenviar-correo-activacion/', data=b'{}', content_type='application/json')
        resp = reenviar_correo_activacion(req)
        self.assertEqual(resp.status_code, 400)

    def test_usuario_no_encontrado(self):
        from unittest.mock import patch, MagicMock
        from firebase_admin import auth as fb_auth
        from .auth_views import reenviar_correo_activacion

        mock_app = MagicMock()
        with patch('smartbids.auth_views._firebase_admin_app', return_value=mock_app), \
             patch('firebase_admin.auth.get_user_by_email', side_effect=fb_auth.UserNotFoundError('No user')):
            req = self.rf.post(
                '/api/reenviar-correo-activacion/',
                data=b'{"email": "inexistente@smartbids.cl"}',
                content_type='application/json'
            )
            resp = reenviar_correo_activacion(req)
            self.assertEqual(resp.status_code, 404)

    def test_usuario_ya_activado(self):
        from unittest.mock import patch, MagicMock
        from .auth_views import reenviar_correo_activacion

        mock_user = MagicMock(email_verified=True)
        with patch('smartbids.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', return_value=mock_user):
            req = self.rf.post(
                '/api/reenviar-correo-activacion/',
                data=b'{"email": "activo@smartbids.cl"}',
                content_type='application/json'
            )
            resp = reenviar_correo_activacion(req)
            self.assertEqual(resp.status_code, 400)

    def test_reenvio_exitoso_y_rate_limit(self):
        from unittest.mock import patch, MagicMock
        from .auth_views import reenviar_correo_activacion

        mock_user = MagicMock(email_verified=False)
        with patch('smartbids.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', return_value=mock_user), \
             patch('smartbids.auth_views._enviar_email_activacion') as mock_envio:
            req = self.rf.post(
                '/api/reenviar-correo-activacion/',
                data=b'{"email": "pendiente@smartbids.cl"}',
                content_type='application/json'
            )
            resp = reenviar_correo_activacion(req)
            self.assertEqual(resp.status_code, 200)
            mock_envio.assert_called_once_with('pendiente@smartbids.cl', app=mock_envio.call_args.kwargs.get('app'), request=req)

            # Segundo intento inmediato debe devolver 429
            req2 = self.rf.post(
                '/api/reenviar-correo-activacion/',
                data=b'{"email": "pendiente@smartbids.cl"}',
                content_type='application/json'
            )
            resp2 = reenviar_correo_activacion(req2)
            self.assertEqual(resp2.status_code, 429)

    def test_correo_malformado_admin(self):
        import json
        from .auth_views import reenviar_correo_activacion

        req = self.rf.post(
            '/api/reenviar-correo-activacion/',
            data=b'{"email": "admin"}',
            content_type='application/json'
        )
        resp = reenviar_correo_activacion(req)
        self.assertEqual(resp.status_code, 400)
        data = json.loads(resp.content.decode('utf-8'))
        self.assertEqual(data.get('status'), 'error')
        self.assertIn('formato del correo ingresado no es válido', data.get('mensaje', ''))
        self.assertNotIn('Malformed email address', data.get('mensaje', ''))

    def test_correo_malformado_en_enviar_codigo_login(self):
        import json
        from .auth_views import enviar_codigo_login

        req = self.rf.post(
            '/api/enviar-codigo-login/',
            data=b'{"email": "admin"}',
            content_type='application/json'
        )
        resp = enviar_codigo_login(req)
        self.assertEqual(resp.status_code, 400)
        data = json.loads(resp.content.decode('utf-8'))
        self.assertEqual(data.get('status'), 'error')
        self.assertIn('formato del correo ingresado no es válido', data.get('mensaje', ''))


class TemplateIngresoTests(SimpleTestCase):
    def test_renderizado_boton_ayuda_y_panel_opciones(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/ingreso/')

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('registro/ingreso.html', request=req)

        # Botón de ayuda en la posición de login actions
        self.assertIn('id="btn-ayuda-login"', html)
        self.assertIn('Ayuda', html)

        # Panel de ayuda con ambas opciones
        self.assertIn('id="ayuda-login-panel"', html)
        self.assertIn('id="forgot-password-link"', html)
        self.assertIn('Recuperar contraseña', html)
        self.assertIn('id="btn-reenviar-activacion"', html)
        self.assertIn('Reenviar activación', html)


class AdministracionGeneralTests(SimpleTestCase):
    def test_navbar_no_contiene_administracion_general(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/mis-licitaciones/')

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('smartbids/Mis_licitaciones.html', request=req, context={'licitaciones': [], 'total_licitaciones': 0})

        self.assertNotIn('id="admin-general-nav"', html)
        self.assertNotIn('Administración General</a>', html)

    def test_perfil_contiene_seccion_admin_en_sesiones(self):
        from unittest.mock import patch
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        rf = RequestFactory()
        req = rf.get('/perfil/')

        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            html = render_to_string('smartbids/perfil.html', request=req)

        self.assertIn('id="perfil-admin-section"', html)
        self.assertIn('Administración General', html)

    def test_acceso_denegado_no_administrador(self):
        from unittest.mock import patch, MagicMock
        from django.test import RequestFactory
        from .page_views import administracion_general_view

        rf = RequestFactory()
        req = rf.get('/administracion-general/')

        mock_susc = MagicMock()
        mock_susc.codigo_estado = MagicMock(codigo_estado=2, nombre_estado='Habilitado')

        with patch('smartbids.page_views.resolver_suscriptor', return_value=mock_susc):
            resp = administracion_general_view(req)
            self.assertEqual(resp.status_code, 302)
            self.assertEqual(resp.url, '/mis-licitaciones/')


class AccionCuentaTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        self.rf = RequestFactory()

    def test_construir_enlace_accion_web(self):
        from .auth_views import _construir_enlace_accion_web
        raw_link = "https://smartbids-auth-users.firebaseapp.com/__/auth/action?mode=verifyEmail&oobCode=XYZ123&apiKey=KEY789&lang=es-419"
        req = self.rf.get('/api/algo/')

        link = _construir_enlace_accion_web(raw_link, request=req, mode_override='verifyEmail')
        self.assertIn('/auth/action/?', link)
        self.assertIn('mode=verifyEmail', link)
        self.assertIn('oobCode=XYZ123', link)
        self.assertIn('apiKey=KEY789', link)

    def test_solicitar_recuperacion_password_validacion(self):
        from .auth_views import solicitar_recuperacion_password
        # Sin email
        req = self.rf.post('/api/solicitar-recuperacion-password/', data=b'{}', content_type='application/json')
        resp = solicitar_recuperacion_password(req)
        self.assertEqual(resp.status_code, 400)

        # Formato inválido
        req_bad = self.rf.post('/api/solicitar-recuperacion-password/', data=b'{"email": "invalido"}', content_type='application/json')
        resp_bad = solicitar_recuperacion_password(req_bad)
        self.assertEqual(resp_bad.status_code, 400)

    def test_solicitar_recuperacion_password_exito(self):
        from unittest.mock import patch, MagicMock
        from .auth_views import solicitar_recuperacion_password
        from django.core.cache import cache

        cache.clear()
        mock_user = MagicMock(email='usuario@smartbids.cl')
        with patch('smartbids.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', return_value=mock_user), \
             patch('firebase_admin.auth.generate_password_reset_link', return_value='https://firebase.com?oobCode=RESET123&apiKey=KEY'), \
             patch('smartbids.auth_views.send_mail') as mock_mail:

            req = self.rf.post(
                '/api/solicitar-recuperacion-password/',
                data=b'{"email": "usuario@smartbids.cl"}',
                content_type='application/json'
            )
            resp = solicitar_recuperacion_password(req)
            self.assertEqual(resp.status_code, 200)
            mock_mail.assert_called_once()
            call_kwargs = mock_mail.call_args.kwargs
            self.assertIn('/auth/action/?', call_kwargs.get('html_message', ''))
            self.assertIn('mode=resetPassword', call_kwargs.get('html_message', ''))

    def test_renderizado_vista_accion_cuenta(self):
        from unittest.mock import patch
        from .page_views import accion_cuenta_view
        req = self.rf.get('/auth/action/?mode=verifyEmail&oobCode=TEST1234')
        with patch('smartbids.context_processors.InfoContacto.objects.first', return_value=None):
            resp = accion_cuenta_view(req)
        self.assertEqual(resp.status_code, 200)
        content = resp.content.decode('utf-8')
        self.assertIn('btn-activar-cuenta', content)
        self.assertIn('reset-password-form', content)


class ValidacionRegistroEmailTests(SimpleTestCase):
    def setUp(self):
        from django.test import RequestFactory
        self.rf = RequestFactory()

    def test_verificar_existencia_gmail_formato(self):
        from .auth_views import verificar_existencia_gmail
        ok, err = verificar_existencia_gmail('abc@gmail.com')
        self.assertFalse(ok)
        self.assertIn('al menos 6 caracteres', err)

        ok2, err2 = verificar_existencia_gmail('usuarioextremadamentelargoconmasde30caracteres@gmail.com')
        self.assertFalse(ok2)
        self.assertIn('más de 30 caracteres', err2)

        ok3, err3 = verificar_existencia_gmail('contacto@miempresa.cl')
        self.assertTrue(ok3)
        self.assertIsNone(err3)

    def test_validar_email_registro_endpoints(self):
        from unittest.mock import patch, MagicMock
        import firebase_admin.auth
        from .auth_views import validar_email_registro

        req = self.rf.post('/api/validar-email-registro/', data=b'{}', content_type='application/json')
        resp = validar_email_registro(req)
        self.assertEqual(resp.status_code, 400)

        with patch('smartbids.auth_views.verificar_existencia_gmail', return_value=(False, 'La cuenta no existe en Gmail')):
            req_gmail = self.rf.post('/api/validar-email-registro/', data=b'{"email": "inexistente@gmail.com"}', content_type='application/json')
            resp_gmail = validar_email_registro(req_gmail)
            self.assertEqual(resp_gmail.status_code, 400)
            self.assertIn('no existe en Gmail', resp_gmail.content.decode('utf-8'))

        mock_user = MagicMock(email='existente@smartbids.cl')
        with patch('smartbids.auth_views.verificar_existencia_gmail', return_value=(True, None)), \
             patch('smartbids.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', return_value=mock_user):
            req_duplicado = self.rf.post('/api/validar-email-registro/', data=b'{"email": "existente@smartbids.cl"}', content_type='application/json')
            resp_duplicado = validar_email_registro(req_duplicado)
            self.assertEqual(resp_duplicado.status_code, 400)
            self.assertIn('Ya existe una cuenta registrada', resp_duplicado.content.decode('utf-8'))

        with patch('smartbids.auth_views.verificar_existencia_gmail', return_value=(True, None)), \
             patch('smartbids.auth_views._firebase_admin_app', return_value=MagicMock()), \
             patch('firebase_admin.auth.get_user_by_email', side_effect=firebase_admin.auth.UserNotFoundError('No user')):
            req_nuevo = self.rf.post('/api/validar-email-registro/', data=b'{"email": "nuevo@smartbids.cl"}', content_type='application/json')
            resp_nuevo = validar_email_registro(req_nuevo)
            self.assertEqual(resp_nuevo.status_code, 200)

    def test_filtrar_licitaciones_sin_perfil_match_cero(self):
        from unittest.mock import patch, MagicMock
        from .licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs

        with patch('smartbids.licitaciones_views.Licitacion.objects', mock_qs), \
             patch('smartbids.licitaciones_views.Preferencia.objects.filter') as mock_pref:
            # Suscriptor sin preferencias configuradas (None o listas vacías)
            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=[],
                pref_tipo_licitacion=[],
                pref_ucom=[],
                pref_palabras_claves=[],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/')
            qs, match_perfil = filtrar_licitaciones(req, suscriptor=MagicMock())

            # El match debe ser 0 y mock_qs no debe haberse filtrado por q_perfil
            self.assertEqual(match_perfil, 0)
            mock_qs.filter.assert_not_called()

    def test_filtrar_licitaciones_con_perfil_aplica_filtro_y_aumenta_match(self):
        from unittest.mock import patch, MagicMock
        from .licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_filtered_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs
        mock_qs.filter.return_value.distinct.return_value = mock_filtered_qs

        with patch('smartbids.licitaciones_views.Licitacion.objects') as mock_licitacion_objects, \
             patch('smartbids.licitaciones_views.Preferencia.objects.filter') as mock_pref:
            mock_licitacion_objects.select_related.return_value.all.return_value.order_by.return_value = mock_qs
            mock_licitacion_objects.filter.return_value.distinct.return_value.count.return_value = 42

            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=['13101'],
                pref_tipo_licitacion=['LE'],
                pref_ucom=[],
                pref_palabras_claves=['servidor', 'computador'],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/?tab=claves')
            qs, match_perfil = filtrar_licitaciones(req, suscriptor=MagicMock())

            # El match con perfil debe ser 42 y el queryset devuelto debe estar filtrado en pestaña de claves
            self.assertEqual(match_perfil, 42)
            self.assertEqual(qs, mock_filtered_qs)

    def test_filtrar_licitaciones_tab_todas_sin_filtro_de_claves(self):
        from unittest.mock import patch, MagicMock
        from .licitaciones_views import filtrar_licitaciones

        mock_qs = MagicMock()
        mock_qs.select_related.return_value.all.return_value.order_by.return_value = mock_qs

        with patch('smartbids.licitaciones_views.Licitacion.objects') as mock_licitacion_objects, \
             patch('smartbids.licitaciones_views.Preferencia.objects.filter') as mock_pref:
            mock_licitacion_objects.select_related.return_value.all.return_value.order_by.return_value = mock_qs
            mock_licitacion_objects.filter.return_value.distinct.return_value.count.return_value = 15

            mock_pref.return_value.first.return_value = MagicMock(
                pref_comunas=[],
                pref_tipo_licitacion=[],
                pref_ucom=[],
                pref_palabras_claves=['software'],
                pref_productos=[]
            )
            req = self.rf.get('/mis-licitaciones/?tab=todas')
            qs, match_perfil = filtrar_licitaciones(req, suscriptor=MagicMock())

            # En pestaña Todas, no se aplica el filtro de palabras clave a qs, pero se calcula el match
            self.assertEqual(match_perfil, 15)
            self.assertEqual(qs, mock_qs)

    def test_mis_licitaciones_view_kpi_match_cero_cuando_perfil_vacio(self):
        from unittest.mock import patch, MagicMock
        from .licitaciones_views import mis_licitaciones_view

        with patch('smartbids.licitaciones_views.resolver_suscriptor', return_value=None), \
             patch('smartbids.licitaciones_views.cache.get_or_set', return_value=62650), \
             patch('smartbids.licitaciones_views.filtrar_licitaciones', return_value=(MagicMock(), 0)), \
             patch('smartbids.licitaciones_views.render') as mock_render, \
             patch('smartbids.licitaciones_views.Paginator') as mock_paginator:
            
            mock_paginator.return_value.get_page.return_value = []
            req = self.rf.get('/mis-licitaciones/')
            mis_licitaciones_view(req)

            self.assertTrue(mock_render.called)
            context = mock_render.call_args[0][2]
            self.assertEqual(context['match_perfil'], 0)



